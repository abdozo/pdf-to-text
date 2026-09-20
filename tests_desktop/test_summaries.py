from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from docx import Document

from waraq.database import Library
from waraq.gemini_client import GeminiClient, GeminiFailure
from waraq.models import ExtractedPage
from waraq.secrets import MemorySecretStore
from waraq.summaries import (
    QUADRANTS, SUMMARY_PROMPT, SUMMARY_SCHEMA, SummaryDocument,
    export_summary, review_input,
)
from waraq.tasks import ConversionRunner


def document(pages, title="موضوع موثق"):
    item = {"label": "فكرة", "text": "النص المستخلص من المصدر", "pdf_pages": [pages[0]]}
    return SummaryDocument.model_validate({
        "source_pages": pages,
        "sheets": [{"title": title, **{key: [item] for key, _title in QUADRANTS}}],
    })


@pytest.fixture
def setup(tmp_path):
    images = [Image.new("RGB", (100, 120), "white") for _ in range(5)]
    path = tmp_path / "source.pdf"
    images[0].save(path, "PDF", save_all=True, append_images=images[1:])
    library = Library(tmp_path / "library")
    book = library.add_book(path)
    key = library.add_key("test", "ref")
    secrets = MemorySecretStore()
    secrets.set("ref", "test-key")
    return library, book, key, secrets


def task(setup, **kwargs):
    kwargs.setdefault("review_batch_size", 1)
    library, book, key, _secrets = setup
    return library.create_task(
        book_id=book["id"], start_page=1, end_page=5, key_id=key,
        model="gemini-test", prompt_id="default-summary", mode="summary",
        dpi=72, overwrite=False, pages_per_request=2, **kwargs,
    )


def fake_client(raw, finish="STOP"):
    calls = []
    client = object.__new__(GeminiClient)
    client.api_key = "test-key"
    client.client = SimpleNamespace(models=SimpleNamespace(generate_content=lambda **kw: calls.append(kw) or SimpleNamespace(
        text=raw, candidates=[SimpleNamespace(finish_reason=finish)], usage_metadata={"total_token_count": 123},
    )))
    return client, calls


def test_summary_request_is_one_document_with_labeled_pdf_images():
    result = document([11, 12])
    client, calls = fake_client(result.model_dump_json())
    image = Image.new("RGB", (100, 120))
    response = client.summarize([(11, image), (12, image)], model="gemini-3.5-flash-lite", prompt=SUMMARY_PROMPT, source_pages=[11, 12])
    assert len(response.document.sheets) == 1
    parts = calls[0]["contents"][0].parts
    assert parts[1].text == "صورة صفحة PDF رقم 11"
    assert parts[3].text == "صورة صفحة PDF رقم 12"
    assert calls[0]["config"].response_json_schema == SUMMARY_SCHEMA
    assert "صفحة منفردة" in parts[0].text


def test_review_sdk_sends_text_only_and_rejects_new_citations():
    original = document([11, 12])  # The only actual citations are to page 11.
    text, pages = review_input([{"document": original.model_dump()}])
    client, calls = fake_client(original.model_dump_json())
    client.summarize([], model="gemini-3.5-flash-lite", prompt=SUMMARY_PROMPT,
                     source_pages=pages, review_text=text)
    assert all(part.inline_data is None for part in calls[0]["contents"][0].parts)
    assert text in calls[0]["contents"][0].parts[-1].text
    original.sheets[0].critique[0].pdf_pages = [12]
    client, _calls = fake_client(original.model_dump_json())
    with pytest.raises(GeminiFailure, match="إحالة"):
        client.summarize([], model="gemini-3.5-flash-lite", prompt=SUMMARY_PROMPT,
                         source_pages=pages, review_text=text)


@pytest.mark.parametrize("bad", ["missing_quadrant", "outside_citation", "wrong_sources", "no_citation", "truncated"])
def test_summary_rejects_incomplete_or_untraceable_responses(bad):
    data = document([11, 12]).model_dump()
    if bad == "missing_quadrant":
        del data["sheets"][0]["critique"]
    elif bad == "outside_citation":
        data["sheets"][0]["insights"][0]["pdf_pages"] = [3]
    elif bad == "wrong_sources":
        data["source_pages"] = [11]
    elif bad == "no_citation":
        data["sheets"][0]["understanding"][0]["pdf_pages"] = []
    client, _calls = fake_client(json.dumps(data), "MAX_TOKENS" if bad == "truncated" else "STOP")
    with pytest.raises(GeminiFailure):
        client.summarize([], model="gemini-3.5-flash-lite", prompt=SUMMARY_PROMPT, source_pages=[11, 12], review_text="[]")


def test_runner_keeps_ocr_and_summaries_separate_and_resumes_batches(setup, monkeypatch):
    library, book, _key, secrets = setup
    extracted = ExtractedPage(is_blank=False, printed_page="٩٩", lines=[{"type": "p", "md": "نسخ محفوظ"}])
    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())
    before = library.get_page(book["id"], 1)
    task_id = task(setup)
    calls = []

    class Client:
        def __init__(self, _key): pass
        def summarize(self, images, **kwargs):
            calls.append(kwargs)
            result = document(kwargs["source_pages"])
            return SimpleNamespace(document=result, raw=result.model_dump_json(), usage={})

    monkeypatch.setattr("waraq.tasks.GeminiClient", Client)
    monkeypatch.setattr("waraq.tasks.render_page", lambda _path, number, _dpi: number)
    def pause(event):
        if event.get("summary_id"):
            library.request_task_action(task_id, "pause")
    runner = ConversionRunner(library, secrets, pause)
    runner.start(task_id); runner.join(5)
    assert library.task(task_id)["state"] == "paused"
    assert library.task(task_id)["completed"] == 2
    runner = ConversionRunner(Library(library.root), secrets)
    runner.start(task_id); runner.join(5)
    assert library.task(task_id)["state"] == "completed"
    assert [call["source_pages"] for call in calls] == [[1, 2], [3, 4], [5]]
    assert all(call["review_text"] == "" for call in calls)
    assert len(library.list_summaries(book["id"])) == 3
    assert library.get_page(book["id"], 1) == before
    assert library.get_page(book["id"], 2)["state"] == "pending"


def saved(setup, pages):
    library, _book, _key, _secrets = setup
    task_id = task(setup)
    library.prepare_task_run(task_id)
    library.start_task_pages(task_id, pages)
    return library.save_summary(task_id, pages, document(pages))


def test_manual_review_sends_only_selected_summaries_and_preserves_originals(setup, monkeypatch):
    library, book, _key, secrets = setup
    first = saved(setup, [1, 2])
    second = saved(setup, [4, 5])
    task_id = task(setup, summary_ids=[second, first])
    calls = []
    class Client:
        def __init__(self, _key): pass
        def review_summary(self, **kwargs):
            calls.append(kwargs)
            result = document(kwargs["target"]["source_pages"], "نسخة مراجعة")
            return SimpleNamespace(document=result, raw=result.model_dump_json(), usage={})
    monkeypatch.setattr("waraq.tasks.GeminiClient", Client)
    monkeypatch.setattr("waraq.tasks.render_page", lambda *_args: pytest.fail("Review must not render PDF pages"))
    library.page_path(book["id"]).unlink()
    runner = ConversionRunner(library, secrets)
    runner.start(task_id); runner.join(5)
    assert library.task(task_id)["state"] == "completed"
    assert [call["target"]["source_pages"] for call in calls] == [[1, 2], [4, 5]]
    assert calls[0]["previous"] is None
    assert calls[0]["following"]["source_pages"] == [4, 5]
    assert calls[1]["previous"]["sheets"][0]["title"] == "نسخة مراجعة"
    assert calls[1]["following"] is None
    records = library.list_summaries(book["id"])
    assert len(records) == 4
    reviewed = [record for record in records if record["task_id"] == task_id]
    assert [record["input_ids"] for record in reviewed] == [[first], [second]]
    with pytest.raises(ValueError, match="أصلها"):
        task(setup, summary_ids=[first, reviewed[0]["id"]])
    # The whole reviewed series can itself be reviewed again.
    assert task(setup, summary_ids=[record["id"] for record in reviewed])


def test_failed_summary_does_not_mark_source_failed_and_can_retry(setup, monkeypatch):
    library, book, _key, secrets = setup
    task_id = task(setup)
    fail = True
    class Client:
        def __init__(self, _key): pass
        def summarize(self, _images, **kwargs):
            if fail and kwargs["source_pages"] == [1, 2]:
                raise GeminiFailure("bad response")
            result = document(kwargs["source_pages"])
            return SimpleNamespace(document=result, raw=result.model_dump_json(), usage={})
    monkeypatch.setattr("waraq.tasks.GeminiClient", Client)
    monkeypatch.setattr("waraq.tasks.render_page", lambda *_args: None)
    runner = ConversionRunner(library, secrets)
    runner.start(task_id); runner.join(5)
    assert library.task(task_id)["state"] == "completed_with_errors"
    assert library.get_page(book["id"], 1)["state"] == "pending"
    fail = False
    runner.start(task_id); runner.join(5)
    assert library.task(task_id)["state"] == "completed"
    assert len(library.list_summaries(book["id"])) == 3


def test_cancelled_batch_is_not_committed(setup):
    library, book, _key, _secrets = setup
    task_id = task(setup)
    library.prepare_task_run(task_id)
    library.start_task_pages(task_id, [1, 2])
    library.request_task_action(task_id, "cancel")
    with pytest.raises(RuntimeError, match="أُلغيت"):
        library.save_summary(task_id, [1, 2], document([1, 2]))
    assert library.list_summaries(book["id"]) == []


def test_missing_summary_is_rejected_before_task_creation(setup):
    library, _book, _key, _secrets = setup
    before = len(library.list_tasks())
    with pytest.raises(ValueError):
        task(setup, summary_ids=["nonexistent"])
    assert len(library.list_tasks()) == before


def test_export_quadrants_and_references_without_losing_text(tmp_path):
    result = document([52, 53], "<script>عنوان</script>")
    result.sheets[0].understanding[0].text = "فكرة <script> ليست شفرة"
    for suffix in ("html", "md", "docx"):
        path = tmp_path / f"summary.{suffix}"
        export_summary(result, path)
        if suffix == "docx":
            output = Document(path)
            assert len(output.tables) == 1
            assert len(output.tables[0].rows) == 2
            assert "[PDF ص 52]" in output.tables[0].cell(0, 0).text
            assert "التلخيص الاستيعابي" in output.tables[0].cell(0, 0).text
        else:
            text = path.read_text()
            assert "[PDF ص 52]" in text
            assert all(title in text for _key, title in QUADRANTS)
            if suffix == "html":
                assert "<script>" not in text
                assert "grid-template-columns:1fr 1fr" in text


def test_legacy_summary_task_keeps_its_original_page_contract(setup, monkeypatch):
    library, book, _key, secrets = setup
    task_id = task(setup)
    with library.transaction() as db:
        db.execute("UPDATE tasks SET summary_contract=0 WHERE id=?", (task_id,))
    calls = []
    class Client:
        def __init__(self, _key): pass
        def extract_pages(self, images, **_kwargs):
            calls.append([number for number, _image in images])
            return SimpleNamespace(pages={number: ExtractedPage(is_blank=False, lines=[{"type":"p","md":"ملخص قديم"}]) for number, _image in images}, raw="{}", usage={})
        def extract(self, _image, **_kwargs):
            return SimpleNamespace(page=ExtractedPage(is_blank=False, lines=[{"type":"p","md":"ملخص قديم"}]), raw="{}", usage={})
    monkeypatch.setattr("waraq.tasks.GeminiClient", Client)
    monkeypatch.setattr("waraq.tasks.render_page", lambda *_args: None)
    runner = ConversionRunner(library, secrets)
    runner.start(task_id); runner.join(5)
    assert library.task(task_id)["state"] == "completed"
    assert calls == [[1, 2], [3, 4]]
    assert library.list_summaries(book["id"]) == []
    assert library.get_page(book["id"], 1)["state"] == "done"
