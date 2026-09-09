from __future__ import annotations

import io
import json
import os
import re
import sqlite3
import threading
import time
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("QT_QUICK_CONTROLS_STYLE", "Basic")
os.environ.setdefault("QSG_RHI_BACKEND", "software")

from PySide6.QtGui import QGuiApplication

from waraq.backup import create_backup, restore_backup
from waraq.bridge import Bridge, local_path
from waraq import database as database_module
from waraq.database import Library, RequestThrottleDelay, safe_rpm
from waraq.exporter import _office_path, export_html, export_markdown, export_word
from waraq.models import ExtractedPage
from waraq.notifications import SYSTEM_SOUNDS
from waraq.secrets import MemorySecretStore
from waraq.tasks import ConversionRunner


def make_pdf(path: Path, count: int = 3) -> Path:
    images = [Image.new("RGB", (420, 594), "white") for _ in range(count)]
    images[0].save(path, "PDF", save_all=True, append_images=images[1:])
    return path


@pytest.fixture
def library(tmp_path: Path) -> Library:
    return Library(tmp_path / "new-desktop-data")


def result(text: str = "نص عربي", *, column: int = 0, handwritten: bool = False) -> ExtractedPage:
    del column
    content = f"*{text}*" if handwritten else text
    return ExtractedPage(
        is_blank=False,
        printed_page="١٢",
        content_markdown=[content],
    )


def test_completion_sound_setting_is_saved_and_used(library: Library) -> None:
    played: list[str] = []
    sound_id = SYSTEM_SOUNDS[1]["id"]
    bridge = Bridge(library, MemorySecretStore(), sound_player=played.append)

    bridge.setNotificationSound(sound_id)
    bridge._apply_task_event(
        {"type": "task", "state": "completed", "task_id": "finished-task"}
    )

    assert played == [sound_id]
    assert library.setting("completion_sound") == {
        "enabled": True,
        "sound_id": sound_id,
    }

    for state in ("failed", "cancelled", "completed_with_errors"):
        bridge._apply_task_event(
            {"type": "task", "state": state, "task_id": f"task-{state}"}
        )

    assert played == [sound_id] * 4

    bridge.setNotificationSoundEnabled(False)
    bridge._apply_task_event(
        {
            "type": "task",
            "state": "completed_with_errors",
            "task_id": "finished-with-errors",
        }
    )

    assert played == [sound_id] * 4
    restored = Bridge(library, MemorySecretStore(), sound_player=played.append)
    assert restored.notificationSoundEnabled is False
    assert restored.notificationSoundId == sound_id


def test_library_references_pdf_without_copying_it(library: Library, tmp_path: Path):
    source = make_pdf(tmp_path / "كتاب.pdf", 4)
    source_bytes = source.read_bytes()
    pdfs_before_import = {path.resolve() for path in tmp_path.rglob("*.pdf")}
    book = library.add_book(source)
    assert book["page_count"] == 4
    assert [page["number"] for page in book["pages"]] == [1, 2, 3, 4]
    assert Path(book["source_path"]) == source.resolve()
    assert source.read_bytes() == source_bytes
    assert {path.resolve() for path in tmp_path.rglob("*.pdf")} == pdfs_before_import
    assert not (library.root / "books").exists()
    printed_prompt = next(item for item in library.list_prompts() if item["id"] == "default-printed")
    assert "كل عنصر في lines يمثل صفًا نصيًا مطبوعًا واحدًا" in printed_prompt["instructions"]
    assert "كتابًا مصورًا إلى كتاب مكتوب مطابق للمطبوع" in printed_prompt["instructions"]
    assert "h2 لأكبر سطر عرض" in printed_prompt["instructions"]
    assert "لا تضع # أو > أو ---" in printed_prompt["instructions"]
    assert '"type":"hr","align":"center","md":""' in printed_prompt["instructions"]
    assert "كل عنوان ظاهر في الوسط يجب أن يحمل align بقيمة center" in printed_prompt["instructions"]
    assert "**العريض**" in printed_prompt["instructions"]


def test_relinking_a_missing_pdf_preserves_existing_conversion(
    library: Library, tmp_path: Path,
) -> None:
    original = make_pdf(tmp_path / "original.pdf", 2)
    book = library.add_book(original, "كتاب محفوظ")
    extracted = result("نص محفوظ")
    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())
    original.unlink()
    replacement = make_pdf(tmp_path / "replacement.pdf", 2)

    assert library.get_book(book["id"])["source_available"] is False

    relinked = library.relink_book_pdf(book["id"], replacement)

    assert Path(relinked["source_path"]) == replacement.resolve()
    assert relinked["source_available"] is True
    assert relinked["name"] == "كتاب محفوظ"
    assert library.page_path(book["id"]) == replacement.resolve()
    assert library.get_page(book["id"], 1)["content_html"] == "<p>نص محفوظ</p>"


def test_relinking_rejects_a_pdf_with_a_different_page_count(
    library: Library, tmp_path: Path,
) -> None:
    original = make_pdf(tmp_path / "original.pdf", 2)
    book = library.add_book(original)
    replacement = make_pdf(tmp_path / "different.pdf", 3)

    with pytest.raises(ValueError, match="عدد صفحات"):
        library.relink_book_pdf(book["id"], replacement)

    assert Path(library.get_book(book["id"])["source_path"]) == original.resolve()


def test_book_can_be_renamed_and_deleted_from_the_library(
    library: Library, tmp_path: Path,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 1))
    key = library.add_key("مفتاح", "delete-secret-ref")
    library.create_task(
        book_id=book["id"], start_page=1, end_page=1,
        key_id=key, model="gemini-test", prompt_id="default-printed",
        mode="printed", dpi=300, overwrite=False,
    )

    renamed = library.rename_book(book["id"], "  الاسم الجديد  ")
    assert renamed["name"] == "الاسم الجديد"

    library.delete_book(book["id"])

    assert library.list_books() == []
    assert library.list_tasks() == []
    with pytest.raises(KeyError, match="الكتاب غير موجود"):
        library.get_book(book["id"])


def test_page_image_cache_follows_a_relinked_pdf(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from PySide6.QtCore import QSize

    from waraq.image_provider import PageImageProvider

    QGuiApplication.instance() or QGuiApplication([])
    first = make_pdf(tmp_path / "first.pdf", 1)
    second = make_pdf(tmp_path / "second.pdf", 1)
    book = library.add_book(first)

    monkeypatch.setattr(
        "waraq.image_provider.render_page",
        lambda path, _number, dpi: Image.new(
            "RGB", (10, 10), "red" if Path(path) == first.resolve() else "blue"
        ),
    )
    provider = PageImageProvider(library)
    before = provider.requestPixmap(f"{book['id']}/1?dpi=180", QSize(), QSize())

    library.relink_book_pdf(book["id"], second)
    after = provider.requestPixmap(f"{book['id']}/1?dpi=180", QSize(), QSize())

    assert before.toImage().pixelColor(0, 0).name() == "#ff0000"
    assert after.toImage().pixelColor(0, 0).name() == "#0000ff"


def test_extraction_editing_and_reapproval_preserve_versions(library: Library, tmp_path: Path):
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 1))
    library.apply_extraction(book["id"], 1, result(), json.dumps(result().model_dump(), ensure_ascii=False))
    page = library.get_page(book["id"], 1)
    page = library.save_regions(book["id"], 1, page["regions"], reviewed=True)
    assert page["reviewed"] == 1
    page["regions"][0]["lines"][0]["text"] = "نص صححه المستخدم"
    edited = library.save_regions(book["id"], 1, page["regions"])
    assert edited["reviewed"] == 0
    assert edited["regions"][0]["lines"][0]["text"] == "نص صححه المستخدم"
    assert library.page_versions(book["id"], 1)


def test_page_can_reset_only_to_its_latest_ai_extraction(
    library: Library, tmp_path: Path,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "reset-source.pdf", 2))
    first_extraction = result("التحويل الأول")
    second_page_extraction = result("الصفحة الثانية")
    library.apply_extraction(
        book["id"], 1, first_extraction, first_extraction.model_dump_json()
    )
    library.apply_extraction(
        book["id"], 2, second_page_extraction, second_page_extraction.model_dump_json()
    )
    library.save_page_content(
        book["id"], 1, "<p>تعديل المستخدم</p>", reviewed=True
    )
    library.save_page_content(
        book["id"], 2, "<p>تعديل الصفحة الثانية</p>"
    )

    reset = library.reset_page_to_extraction(book["id"], 1)

    assert reset["content_html"] == "<p>التحويل الأول</p>"
    assert reset["reviewed"] == 0
    assert reset["can_reset_to_extraction"] is True
    assert library.get_page(book["id"], 2)["content_html"] == (
        "<p>تعديل الصفحة الثانية</p>"
    )
    assert "reset_to_extraction" in {
        version["reason"] for version in library.page_versions(book["id"], 1)
    }


def test_new_extraction_replaces_the_page_reset_baseline(
    library: Library, tmp_path: Path,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "new-baseline.pdf", 1))
    first = result("تحويل قديم")
    latest = result("تحويل جديد")
    library.apply_extraction(book["id"], 1, first, first.model_dump_json())
    library.save_page_content(book["id"], 1, "<p>تعديل أول</p>")
    library.apply_extraction(book["id"], 1, latest, latest.model_dump_json())
    library.save_page_content(book["id"], 1, "<p>تعديل بعد التحويل الجديد</p>")

    reset = library.reset_page_to_extraction(book["id"], 1)

    assert reset["content_html"] == "<p>تحويل جديد</p>"


def test_existing_page_can_recover_its_reset_baseline_from_raw_result(
    library: Library, tmp_path: Path,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "existing-page.pdf", 1))
    extracted = result("النص القادم من الذكاء الاصطناعي")
    library.apply_extraction(
        book["id"], 1, extracted, extracted.model_dump_json()
    )
    library.save_page_content(book["id"], 1, "<p>تعديل محفوظ تلقائيًا</p>")
    with library.transaction() as database:
        database.execute(
            """UPDATE pages
               SET original_content_html='', original_content_markdown=''
               WHERE book_id=? AND number=1""",
            (book["id"],),
        )

    reset = library.reset_page_to_extraction(book["id"], 1)

    assert reset["content_html"] == "<p>النص القادم من الذكاء الاصطناعي</p>"
    assert reset["original_content_html"] == reset["content_html"]


def test_extraction_markdown_reaches_the_editor_with_one_block_per_visual_line(
    library: Library, tmp_path: Path,
):
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 1))
    extracted = ExtractedPage(
        is_blank=False,
        printed_page="٧",
        lines=[
            {"type": "h2", "md": "عنوان الباب"},
            {"type": "p", "md": "نص **عريض** (١)"},
            {"type": "p", "md": "1. البند الأول"},
            {"type": "p", "md": "2. البند الثاني"},
        ],
    )

    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())
    page = library.get_page(book["id"], 1)

    assert page["printed_page"] == "٧"
    assert page["content_markdown"] == (
        "## عنوان الباب\n\nنص **عريض** (١)\n\n"
        "1. البند الأول\n\n2. البند الثاني"
    )
    assert '<h2 class="ql-align-center">عنوان الباب</h2>' in page["content_html"]
    assert "<strong>عريض</strong> (١)" in page["content_html"]
    assert page["content_html"].count("<ol>") == 1
    assert page["content_html"].count('data-list="ordered"') == 2
    assert [line["text"] for line in page["regions"][0]["lines"]] == [
        "عنوان الباب", "نص عريض (١)", "البند الأول", "البند الثاني",
    ]


def test_extraction_preserves_small_footnotes_and_horizontal_separator(
    library: Library, tmp_path: Path,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "footnote-source.pdf", 1))
    extracted = ExtractedPage(
        is_blank=False,
        printed_page="١",
        lines=[
            {"type": "p", "md": "المتن"},
            {"type": "hr", "md": ""},
            {"type": "note", "md": "(١) نص الحاشية"},
        ],
    )

    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())

    assert library.get_page(book["id"], 1)["content_html"] == (
        "<p>المتن</p>"
        "<hr>"
        '<blockquote class="ql-size-small"><p>(١) نص الحاشية</p></blockquote>'
    )


def test_extraction_preserves_alignment_independently_from_line_type(
    library: Library, tmp_path: Path,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "aligned-source.pdf", 1))
    extracted = ExtractedPage(
        is_blank=False,
        printed_page=None,
        lines=[
            {"type": "p", "align": "center", "md": "سطر عادي متوسط"},
            {"type": "h3", "align": "right", "md": "عنوان يميني"},
            {"type": "p", "align": "left", "md": "Latin line"},
            {"type": "p", "align": "justify", "md": "متن مضبوط"},
        ],
    )

    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())
    page = library.get_page(book["id"], 1)

    assert page["printed_page"] == ""
    assert '<p class="ql-align-center">سطر عادي متوسط</p>' in page["content_html"]
    assert '<h3 class="ql-align-right">عنوان يميني</h3>' in page["content_html"]
    assert '<p class="ql-align-left">Latin line</p>' in page["content_html"]
    assert '<p class="ql-align-justify">متن مضبوط</p>' in page["content_html"]

    saved = library.save_page_content(
        book["id"],
        1,
        page["content_html"],
        content_markdown=page["content_markdown"],
    )
    assert '<p class="ql-align-center">سطر عادي متوسط</p>' in saved["content_html"]
    assert '<h3 class="ql-align-right">عنوان يميني</h3>' in saved["content_html"]


def test_rich_html_converts_non_breaking_spaces_to_wrappable_spaces(
    library: Library, tmp_path: Path,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 1))

    page = library.save_page_content(
        book["id"], 1, "<p>نص\u00a0عربي\u00a0يجب\u00a0أن\u00a0يلتف</p>"
    )

    assert "\u00a0" not in page["content_html"]
    assert page["content_html"] == "<p>نص عربي يجب أن يلتف</p>"


def test_task_snapshot_and_resume_skip_completed(library: Library, tmp_path: Path):
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 2))
    key = library.add_key("أساسي", "secret-ref")
    library.apply_extraction(book["id"], 1, result(), "{}")
    task_id = library.create_task(book_id=book["id"], start_page=1, end_page=2, key_id=key,
                                  model="gemini-3.5-flash-lite", prompt_id="default-printed", mode="printed", dpi=300, overwrite=False)
    task = library.task(task_id)
    assert library.page_numbers_for_task(task) == [2]
    assert task["prompt_snapshot"]
    overwrite_id = library.create_task(book_id=book["id"], start_page=1, end_page=2, key_id=key,
                                       model="gemini-3.5-flash-lite", prompt_id="default-printed", mode="printed", dpi=300, overwrite=True)
    assert library.page_numbers_for_task(library.task(overwrite_id)) == [1, 2]


def test_named_key_owns_its_internal_quota_group(library: Library) -> None:
    key_id = library.add_key("المفتاح الأساسي", "secret-ref")

    keys = library.list_keys()
    assert [(key["id"], key["name"]) for key in keys] == [
        (key_id, "المفتاح الأساسي")
    ]
    assert any(
        limit["key_id"] == key_id
        and limit["key_name"] == "المفتاح الأساسي"
        for limit in library.quota_snapshot()
    )


def test_deleting_named_key_removes_its_metadata_and_quota_group(
    library: Library,
) -> None:
    key_id = library.add_key("مفتاح للحذف", "secret-ref")
    deleted = library.delete_key(key_id)

    assert deleted["secret_ref"] == "secret-ref"
    assert library.list_keys() == []
    assert not any(
        limit["key_id"] == key_id for limit in library.quota_snapshot()
    )
    with library.connect() as db:
        assert not db.execute(
            "SELECT 1 FROM quota_limits WHERE key_id=?", (key_id,)
        ).fetchone()
    with pytest.raises(KeyError):
        library.key_metadata(key_id)


def test_bridge_deletes_named_key_from_library_and_secret_store(
    library: Library,
) -> None:
    secrets = MemorySecretStore()
    bridge = Bridge(library, secrets)
    bridge.createKey("مفتاح شخصي", "AIza-secret")
    key = bridge.keys[0]
    metadata = library.key_metadata(key["id"])
    assert secrets.get(metadata["secret_ref"]) == "AIza-secret"

    bridge.deleteKey(key["id"])

    assert bridge.keys == []
    assert metadata["secret_ref"] not in secrets.values


def test_bridge_keeps_key_and_secret_while_a_linked_task_can_resume(
    library: Library, tmp_path: Path,
) -> None:
    secrets = MemorySecretStore()
    bridge = Bridge(library, secrets)
    bridge.createKey("مفتاح مهمة", "AIza-secret")
    key_id = bridge.keys[0]["id"]
    metadata = library.key_metadata(key_id)
    book = library.add_book(make_pdf(tmp_path / "linked-task.pdf", 1))
    library.create_task(
        book_id=book["id"], start_page=1, end_page=1, key_id=key_id,
        model="gemini-test", prompt_id="default-printed", mode="printed",
        dpi=300, overwrite=False,
    )

    bridge.deleteKey(key_id)

    assert library.key_metadata(key_id)["id"] == key_id
    assert secrets.get(metadata["secret_ref"]) == "AIza-secret"
    assert "أكمل أو ألغِ" in bridge.toast


def test_task_persists_the_selected_pages_per_request(library: Library, tmp_path: Path):
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 4))
    key = library.add_key("أساسي", "secret-ref")

    task_id = library.create_task(
        book_id=book["id"], start_page=1, end_page=4,
        key_id=key, model="gemini-3.5-flash-lite", prompt_id="default-printed",
        mode="printed", dpi=300, overwrite=False, pages_per_request=4,
    )

    assert library.task(task_id)["pages_per_request"] == 4


def test_request_start_is_throttled_by_project_and_model(
    library: Library, tmp_path: Path,
) -> None:
    task_id, _secrets = _conversion_task(library, tmp_path, page_count=1)
    task = library.task(task_id)
    task["model"] = "gemini-2.5-flash"

    library.start_request(task, 1)

    with pytest.raises(RequestThrottleDelay) as error:
        library.start_request(task, 1)
    assert error.value.wait_seconds >= 15


def test_request_throttle_is_independent_for_each_named_key(
    library: Library, tmp_path: Path,
) -> None:
    task_id, _secrets = _conversion_task(library, tmp_path, page_count=1)
    task = library.task(task_id)
    task["model"] = "gemini-2.5-flash-lite"
    second_key = library.add_key("مفتاح ثان", "second-secret")

    library.start_request(task, 1)
    task["key_id"] = second_key

    library.start_request(task, 1)


def test_safe_rpm_stays_one_below_google_limit_and_is_persisted(
    library: Library,
) -> None:
    key = library.add_key("مفتاح الحدود", "secret-ref")
    limits = {
        item["model"]: item
        for item in library.quota_snapshot()
        if item["key_id"] == key
    }

    assert safe_rpm(15) == 14
    assert safe_rpm(5) == 4
    assert limits["gemini-3.5-flash-lite"]["rpm"] == 15
    assert limits["gemini-3.5-flash-lite"]["safe_rpm"] == 14
    assert limits["gemini-2.5-flash"]["rpm"] == 5
    assert limits["gemini-2.5-flash"]["safe_rpm"] == 4

    reopened = Library(library.root)
    reopened_limits = {
        item["model"]: item
        for item in reopened.quota_snapshot()
        if item["key_id"] == key
    }
    assert reopened_limits["gemini-3.5-flash-lite"]["safe_rpm"] == 14
    assert reopened_limits["gemini-2.5-flash"]["safe_rpm"] == 4


def test_recent_requests_still_block_the_model_after_restart(
    library: Library, tmp_path: Path,
) -> None:
    task_id, _secrets = _conversion_task(library, tmp_path, page_count=1)
    task = library.task(task_id)
    task["model"] = "gemini-2.5-flash"
    library.start_request(task, 1)

    reopened = Library(library.root)

    with pytest.raises(RequestThrottleDelay) as error:
        reopened.start_request(task, 1)
    assert error.value.wait_seconds >= 15


def test_existing_tasks_are_migrated_to_durable_per_page_progress(tmp_path: Path) -> None:
    migrations = database_module.MIGRATIONS
    root = tmp_path / "legacy-data"
    root.mkdir()
    database_path = root / "library.sqlite3"
    now = database_module.utc_now()
    with sqlite3.connect(database_path) as db:
        db.execute(
            "CREATE TABLE schema_migrations (version INTEGER PRIMARY KEY, applied_at TEXT NOT NULL)"
        )
        for version, migration in enumerate(migrations[:2], 1):
            db.executescript(migration)
            db.execute("INSERT INTO schema_migrations VALUES (?, ?)", (version, now))
        db.execute(
            "INSERT INTO books VALUES ('book', 'كتاب قديم', 'source.pdf', 2, ?, ?, 1, NULL)",
            (now, now),
        )
        db.executemany(
            """INSERT INTO pages
               (id,book_id,number,width,height,state,reviewed,printed_page,is_blank,
                error,raw_result,updated_at,content_html)
               VALUES (?, 'book', ?, 420, 594, ?, 0, '', 0, '', ?, ?, '')""",
            [
                ("page-1", 1, "done", "{}", now),
                ("page-2", 2, "pending", "", now),
            ],
        )
        db.execute("INSERT INTO projects VALUES ('project', 'مشروع قديم', ?)", (now,))
        db.execute(
            "INSERT INTO api_keys VALUES ('key', 'project', 'مفتاح قديم', 'secret-ref', ?, NULL)",
            (now,),
        )
        db.execute(
            "INSERT INTO prompts VALUES ('default-printed', 'قديم', 'printed', 'prompt', 1, ?, ?)",
            (now, now),
        )
        db.execute(
            """INSERT INTO tasks
               (id,book_id,start_page,end_page,state,project_id,key_id,model,prompt_id,
                prompt_snapshot,mode,dpi,overwrite,stop_requested,current_page,completed,
                failed,error,created_at,updated_at)
               VALUES ('legacy-task', ?, 1, 2, 'paused', ?, ?, 'gemini-test',
                       'default-printed', ?, 'printed', 300, 0, 0, NULL, 0, 0, '', ?, ?)""",
            ("book", "project", "key", "prompt", now, now),
        )

    upgraded = Library(root)
    task = upgraded.task("legacy-task")

    assert task["state"] == "paused"
    assert task["total"] == 2
    assert task["skipped"] == 1
    assert task["remaining"] == 1
    assert upgraded.page_numbers_for_task(task) == [2]
    with upgraded.connect() as db:
        assert not db.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='projects'"
        ).fetchone()
        assert "project_id" not in {
            row["name"] for row in db.execute("PRAGMA table_info(api_keys)")
        }
        assert db.execute(
            "SELECT 1 FROM quota_limits WHERE key_id='key'"
        ).fetchone()


def _conversion_task(library: Library, tmp_path: Path, page_count: int = 3) -> tuple[str, MemorySecretStore]:
    book = library.add_book(make_pdf(tmp_path / "task-source.pdf", page_count))
    secrets = MemorySecretStore()
    secrets.set("secret-ref", "test-key")
    key = library.add_key("المفتاح", "secret-ref")
    task_id = library.create_task(
        book_id=book["id"], start_page=1, end_page=page_count,
        key_id=key, model="gemini-test",
        prompt_id="default-printed", mode="printed", dpi=300, overwrite=True,
    )
    return task_id, secrets


def test_conversion_reads_the_pdf_from_its_original_path_without_copying(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    task_id, secrets = _conversion_task(library, tmp_path, page_count=1)
    task = library.task(task_id)
    source = library.page_path(task["book_id"])
    source_bytes = source.read_bytes()
    pdfs_before_conversion = {path.resolve() for path in tmp_path.rglob("*.pdf")}
    rendered_paths: list[Path] = []

    class FakeClient:
        def __init__(self, _key: str):
            pass

        def extract(self, image, **_kwargs):
            extracted = result(f"صفحة {image}")
            return SimpleNamespace(
                page=extracted,
                raw=extracted.model_dump_json(),
                usage={},
            )

    def record_render_path(path: Path, number: int, _dpi: int) -> int:
        rendered_paths.append(Path(path).resolve())
        return number

    monkeypatch.setattr("waraq.tasks.GeminiClient", FakeClient)
    monkeypatch.setattr("waraq.tasks.render_page", record_render_path)

    runner = ConversionRunner(library, secrets)
    runner.start(task_id)
    runner.join(2)

    assert library.task(task_id)["state"] == "completed"
    assert rendered_paths == [source.resolve()]
    assert source.read_bytes() == source_bytes
    assert {path.resolve() for path in tmp_path.rglob("*.pdf")} == pdfs_before_conversion


def test_pause_then_resume_continues_without_reconverting_saved_pages(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    task_id, secrets = _conversion_task(library, tmp_path)
    first_started = threading.Event()
    release_first = threading.Event()
    converted_pages: list[int] = []

    class FakeClient:
        def __init__(self, _key: str):
            pass

        def extract(self, image, **_kwargs):
            converted_pages.append(image)
            if image == 1:
                first_started.set()
                assert release_first.wait(2)
            extracted = result(f"صفحة {image}")
            return SimpleNamespace(page=extracted, raw=extracted.model_dump_json(), usage={})

    monkeypatch.setattr("waraq.tasks.GeminiClient", FakeClient)
    monkeypatch.setattr("waraq.tasks.render_page", lambda _path, number, _dpi: number)
    runner = ConversionRunner(library, secrets)

    runner.start(task_id)
    assert first_started.wait(2)
    runner.request_pause(task_id)
    release_first.set()
    runner.join(2)

    paused = library.task(task_id)
    assert paused["state"] == "paused"
    assert paused["completed"] == 1
    assert paused["remaining"] == 2

    runner.start(task_id)
    runner.join(2)

    completed = library.task(task_id)
    assert completed["state"] == "completed"
    assert completed["completed"] == 3
    assert completed["remaining"] == 0
    assert converted_pages == [1, 2, 3]


def test_conversion_runner_waits_for_a_local_rate_limit_slot(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    task_id, secrets = _conversion_task(library, tmp_path, page_count=1)
    original_start_request = library.start_request
    attempts = 0

    def gated_start_request(task, page_number):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise RequestThrottleDelay(0)
        return original_start_request(task, page_number)

    class FakeClient:
        def __init__(self, _key: str):
            pass

        def extract(self, image, **_kwargs):
            extracted = result(f"صفحة {image}")
            return SimpleNamespace(page=extracted, raw=extracted.model_dump_json(), usage={})

    monkeypatch.setattr(library, "start_request", gated_start_request)
    monkeypatch.setattr("waraq.tasks.GeminiClient", FakeClient)
    monkeypatch.setattr("waraq.tasks.render_page", lambda _path, number, _dpi: number)

    runner = ConversionRunner(library, secrets)
    runner.start(task_id)
    runner.join(2)

    assert library.task(task_id)["state"] == "completed"
    assert attempts == 2


def test_conversion_runner_groups_pages_into_the_selected_request_size(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    task_id, secrets = _conversion_task(library, tmp_path, page_count=5)
    with library.transaction() as db:
        db.execute("UPDATE tasks SET pages_per_request=2 WHERE id=?", (task_id,))
    sent_batches: list[list[int]] = []

    class FakeClient:
        def __init__(self, _key: str):
            pass

        def extract(self, image, **_kwargs):
            sent_batches.append([image])
            extracted = result(f"صفحة {image}")
            return SimpleNamespace(page=extracted, raw=extracted.model_dump_json(), usage={})

        def extract_pages(self, images, **_kwargs):
            numbers = [number for number, _image in images]
            sent_batches.append(numbers)
            pages = {number: result(f"صفحة {number}") for number in numbers}
            return SimpleNamespace(pages=pages, raw=json.dumps({"pages": numbers}), usage={})

    monkeypatch.setattr("waraq.tasks.GeminiClient", FakeClient)
    monkeypatch.setattr("waraq.tasks.render_page", lambda _path, number, _dpi: number)

    runner = ConversionRunner(library, secrets)
    runner.start(task_id)
    runner.join(2)

    assert library.task(task_id)["state"] == "completed"
    assert sent_batches == [[1, 2], [3, 4], [5]]
    requests = list(reversed(library.request_rows()))
    assert [json.loads(item["page_numbers"]) for item in requests] == [[1, 2], [3, 4], [5]]
    assert library.usage_summary()["requests"] == 3


def test_conversion_can_be_paused_while_waiting_for_rate_limit(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    task_id, secrets = _conversion_task(library, tmp_path, page_count=1)
    waiting = threading.Event()

    def gated_start_request(_task, _page_number):
        raise RequestThrottleDelay(10)

    def progress(event):
        if event.get("state") == "waiting":
            waiting.set()

    monkeypatch.setattr(library, "start_request", gated_start_request)
    runner = ConversionRunner(library, secrets, progress)
    runner.start(task_id)
    assert waiting.wait(2)
    runner.request_pause(task_id)
    runner.join(2)

    task = library.task(task_id)
    assert task["state"] == "paused"
    assert task["remaining"] == 1


def test_cancel_preserves_saved_pages_and_makes_task_terminal(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    task_id, secrets = _conversion_task(library, tmp_path)
    second_started = threading.Event()
    release_second = threading.Event()

    class FakeClient:
        def __init__(self, _key: str):
            pass

        def extract(self, image, **_kwargs):
            if image == 2:
                second_started.set()
                assert release_second.wait(2)
            extracted = result(f"صفحة {image}")
            return SimpleNamespace(page=extracted, raw=extracted.model_dump_json(), usage={})

    monkeypatch.setattr("waraq.tasks.GeminiClient", FakeClient)
    monkeypatch.setattr("waraq.tasks.render_page", lambda _path, number, _dpi: number)
    runner = ConversionRunner(library, secrets)

    runner.start(task_id)
    assert second_started.wait(2)
    runner.request_cancel(task_id)
    release_second.set()
    for _ in range(100):
        if not any(
            thread.name.startswith(f"waraq-{task_id[:8]}")
            for thread in threading.enumerate()
        ):
            break
        time.sleep(0.01)

    cancelled = library.task(task_id)
    assert cancelled["state"] == "cancelled"
    assert cancelled["completed"] == 1
    assert cancelled["remaining"] == 2
    assert library.get_page(cancelled["book_id"], 1)["state"] == "done"
    with pytest.raises(RuntimeError):
        runner.start(task_id)


def test_cancel_is_immediate_and_a_new_task_can_start_while_request_returns(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    task_id, secrets = _conversion_task(library, tmp_path, page_count=1)
    task = library.task(task_id)
    first_started = threading.Event()
    release_first = threading.Event()
    close_started = threading.Event()
    release_close = threading.Event()
    calls = 0

    class FakeClient:
        def __init__(self, _key: str):
            pass

        def close(self):
            close_started.set()
            assert release_close.wait(2)

        def extract(self, image, **_kwargs):
            nonlocal calls
            calls += 1
            if calls == 1:
                first_started.set()
                assert release_first.wait(2)
            extracted = result(f"صفحة {image}")
            return SimpleNamespace(
                page=extracted,
                raw=extracted.model_dump_json(),
                usage={},
            )

    monkeypatch.setattr("waraq.tasks.GeminiClient", FakeClient)
    monkeypatch.setattr("waraq.tasks.render_page", lambda _path, number, _dpi: number)
    runner = ConversionRunner(library, secrets)

    runner.start(task_id)
    assert first_started.wait(2)
    cancel_started_at = time.monotonic()
    runner.request_cancel(task_id)
    assert time.monotonic() - cancel_started_at < 0.5
    assert close_started.wait(0.5)

    cancelled = library.task(task_id)
    assert cancelled["state"] == "cancelled"
    assert cancelled["completed"] == 0
    assert runner.running is False

    next_task_id = library.create_task(
        book_id=task["book_id"], start_page=1, end_page=1,
        key_id=task["key_id"], model=task["model"],
        prompt_id=task["prompt_id"], mode=task["mode"], dpi=task["dpi"],
        overwrite=True,
    )
    runner.start(next_task_id)
    runner.join(2)
    assert library.task(next_task_id)["state"] == "completed"

    release_first.set()
    release_close.set()
    for _ in range(100):
        if not any(
            thread.name.startswith(f"waraq-{task_id[:8]}")
            for thread in threading.enumerate()
        ):
            break
        time.sleep(0.01)
    assert library.task(task_id)["state"] == "cancelled"
    assert library.task(task_id)["completed"] == 0


def test_bridge_reloads_persisted_tasks_after_the_application_reopens(
    library: Library, tmp_path: Path,
) -> None:
    task_id, _secrets = _conversion_task(library, tmp_path)
    library.update_task(task_id, state="paused")

    bridge = Bridge(library, MemorySecretStore())

    assert bridge.currentTask["id"] == task_id
    assert bridge.tasks[0]["id"] == task_id
    bridge.go("task")
    assert bridge.screen == "task"


def test_backup_round_trip_keeps_database_without_copying_pdf(library: Library, tmp_path: Path):
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 2), "قبل النسخ")
    backup = create_backup(library, tmp_path / "copy.waraq-backup")
    with ZipFile(backup) as archive:
        assert "manifest.json" in archive.namelist()
        assert "library.sqlite3" in archive.namelist()
        assert not any(name.endswith(".pdf") for name in archive.namelist())
        assert json.loads(archive.read("manifest.json"))["pdf_storage"] == "external"
    library.add_book(make_pdf(tmp_path / "later.pdf", 1), "بعد النسخ")
    restore_backup(library, backup)
    assert [item["name"] for item in library.list_books()] == ["قبل النسخ"]
    assert library.page_path(book["id"]) == (tmp_path / "source.pdf").resolve()
    assert library.page_path(book["id"]).exists()


def test_restoring_a_legacy_backup_never_extracts_its_managed_pdf(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = make_pdf(tmp_path / "original-location.pdf", 1)
    book = library.add_book(source)
    current_backup = create_backup(library, tmp_path / "current.waraq-backup")
    with ZipFile(current_backup) as archive:
        database_snapshot = archive.read("library.sqlite3")

    legacy_backup = tmp_path / "legacy.waraq-backup"
    with ZipFile(legacy_backup, "w") as archive:
        archive.writestr(
            "manifest.json",
            json.dumps({"format": 1, "application": "Warraq"}),
        )
        archive.writestr("library.sqlite3", database_snapshot)
        archive.writestr("books/managed-copy.pdf", source.read_bytes())

    extracted_members: list[str] = []
    original_extract = ZipFile.extract

    def record_extract(archive: ZipFile, member, *args, **kwargs):
        extracted_members.append(member if isinstance(member, str) else member.filename)
        return original_extract(archive, member, *args, **kwargs)

    monkeypatch.setattr(ZipFile, "extract", record_extract)
    monkeypatch.setattr(
        ZipFile,
        "extractall",
        lambda *_args, **_kwargs: pytest.fail("يجب ألا تُفك كل محتويات النسخة الاحتياطية"),
    )
    restore_backup(library, legacy_backup)

    assert extracted_members == ["library.sqlite3"]
    assert library.page_path(book["id"]) == source.resolve()
    assert not (library.root / "books").exists()
    assert {path.resolve() for path in tmp_path.rglob("*.pdf")} == {source.resolve()}


def test_missing_external_pdf_has_a_clear_error(library: Library, tmp_path: Path):
    source = make_pdf(tmp_path / "source.pdf", 1)
    book = library.add_book(source)
    source.unlink()

    with pytest.raises(FileNotFoundError, match="أعد ربط الكتاب"):
        library.page_path(book["id"])


def test_local_file_url_is_converted_by_qt(tmp_path: Path):
    source = tmp_path / "كتاب بمسافات.pdf"
    url = source.as_uri()

    assert local_path(url) == source


def test_word_export_is_editable_without_writing_a_json_report(library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 2))
    for number in (1, 2):
        extracted = result(f"صفحة {number}", column=number - 1)
        library.apply_extraction(book["id"], number, extracted, extracted.model_dump_json())
        page = library.get_page(book["id"], number)
        library.save_regions(book["id"], number, page["regions"], reviewed=True)
    monkeypatch.setattr("waraq.exporter._office_path", lambda: None)
    destination = tmp_path / "book.docx"
    report = export_word(library, book["id"], 1, 2, destination)
    assert destination.exists() and report["expected_pages"] == 2 and not report["verified"]
    with ZipFile(destination) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "صفحة 1" in xml and "صفحة 2" in xml
    assert not destination.with_suffix(".report.json").exists()


def test_word_export_writes_portable_rtl_paragraph_and_run_properties(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "rtl-source.pdf", 1))
    extracted = result("هذا نص عربي لاختبار اتجاه الكتابة")
    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())
    page = library.get_page(book["id"], 1)
    library.save_regions(book["id"], 1, page["regions"], reviewed=True)
    monkeypatch.setattr("waraq.exporter._office_path", lambda: None)

    destination = tmp_path / "rtl.docx"
    export_word(library, book["id"], 1, 1, destination)

    with ZipFile(destination) as archive:
        document = archive.read("word/document.xml").decode("utf-8")
        styles = archive.read("word/styles.xml").decode("utf-8")

    paragraph_properties = re.findall(r"<w:pPr>(.*?)</w:pPr>", document)
    arabic_runs = re.findall(
        r"<w:r>(.*?)<w:t[^>]*>[^<]*[\u0600-\u06ff][^<]*</w:t>.*?</w:r>",
        document,
    )
    assert paragraph_properties
    assert all('<w:bidi w:val="1"/>' in properties for properties in paragraph_properties)
    assert all('<w:jc w:val="start"/>' in properties for properties in paragraph_properties)
    assert all(properties.index("<w:bidi") < properties.index("<w:jc") for properties in paragraph_properties)
    assert arabic_runs
    assert all('<w:rtl w:val="1"/>' in run for run in arabic_runs)
    assert all('<w:cs w:val="1"/>' in run for run in arabic_runs)
    assert all('w:bidi="ar-SA"' in run for run in arabic_runs)
    normal_style = re.search(
        r'<w:style[^>]+w:styleId="Normal".*?</w:style>', styles, re.S
    )
    assert normal_style
    assert '<w:bidi w:val="1"/>' in normal_style.group(0)
    assert '<w:jc w:val="start"/>' in normal_style.group(0)


def test_word_export_renders_unaligned_arabic_at_the_physical_right_edge(
    library: Library, tmp_path: Path,
) -> None:
    if not _office_path():
        pytest.skip("LibreOffice is not available for rendered alignment verification")
    book = library.add_book(make_pdf(tmp_path / "rendered-rtl-source.pdf", 1))
    library.save_page_content(
        book["id"], 1, "<p>سطر قصير لاختبار المحاذاة.</p>", reviewed=True
    )

    destination = tmp_path / "rendered-rtl.docx"
    report = export_word(library, book["id"], 1, 1, destination)
    verification = destination.with_suffix(".verification.pdf")
    if not report["verified"] or not verification.exists():
        pytest.skip("LibreOffice could not render the Word verification file")

    import pypdfium2 as pdfium

    with pdfium.PdfDocument(verification) as pdf:
        page = pdf[0]
        text_page = page.get_textpage()
        extracted = text_page.get_text_range()
        cursor = 0
        target_range = None
        for line in extracted.splitlines(keepends=True):
            length = len(line.rstrip("\r\n"))
            if "المحاذاة" in line:
                target_range = range(cursor, cursor + length)
                break
            cursor += len(line)
        assert target_range is not None
        boxes = [text_page.get_charbox(index) for index in target_range]
        assert min(box[0] for box in boxes) > page.get_width() * 0.45
        assert max(box[2] for box in boxes) > page.get_width() * 0.85


def test_word_export_preserves_rich_editor_paragraph_alignment(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "aligned-source.pdf", 1))
    extracted = result("نص مؤقت")
    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())
    library.save_page_content(
        book["id"],
        1,
        '<p style="text-align: center">سطر في الوسط</p>'
        '<p style="text-align: right">سطر على اليمين</p>'
        '<p style="text-align: left">Left aligned</p>',
        reviewed=True,
    )
    monkeypatch.setattr("waraq.exporter._office_path", lambda: None)

    destination = tmp_path / "aligned.docx"
    export_word(library, book["id"], 1, 1, destination)

    document = Document(destination)
    paragraphs = {
        paragraph.text: paragraph._p.pPr.xml
        for paragraph in document.paragraphs
        if paragraph.text
    }
    assert 'w:jc w:val="center"' in paragraphs["سطر في الوسط"]
    assert 'w:jc w:val="start"' in paragraphs["سطر على اليمين"]
    assert 'w:jc w:val="end"' in paragraphs["Left aligned"]


def test_word_export_renders_horizontal_separator_and_small_footnote(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "footnote-export-source.pdf", 1))
    extracted = ExtractedPage(
        is_blank=False,
        printed_page="١",
        content_markdown=[
            "المتن",
            "---",
            "> (١) نص الحاشية",
        ],
    )
    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())
    library.save_page_content(
        book["id"], 1, library.get_page(book["id"], 1)["content_html"], reviewed=True,
    )
    monkeypatch.setattr("waraq.exporter._office_path", lambda: None)

    destination = tmp_path / "footnote.docx"
    export_word(library, book["id"], 1, 1, destination)

    document = Document(destination)
    note = next(paragraph for paragraph in document.paragraphs if paragraph.text == "(١) نص الحاشية")
    body = next(paragraph for paragraph in document.paragraphs if paragraph.text == "المتن")
    with ZipFile(destination) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")

    assert '<w:pBdr><w:bottom w:val="single"' in xml
    assert note.runs[0].font.size.pt < body.runs[0].font.size.pt


def test_word_export_never_makes_small_notes_larger_than_dense_arabic_body(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "dense-page.pdf", 1))
    body_lines = "".join(
        f"<p>هذا سطر عربي طويل من متن الصفحة رقم {number} ويجب أن يبقى أكبر من الحاشية.</p>"
        for number in range(1, 45)
    )
    library.save_page_content(
        book["id"],
        1,
        body_lines
        + '<hr><blockquote class="ql-size-small"><p>١. حاشية صغيرة</p></blockquote>',
        reviewed=True,
    )
    monkeypatch.setattr("waraq.exporter._office_path", lambda: None)

    destination = tmp_path / "dense-page.docx"
    export_word(library, book["id"], 1, 1, destination)

    document = Document(destination)
    body = next(paragraph for paragraph in document.paragraphs if "متن الصفحة رقم" in paragraph.text)
    note = next(paragraph for paragraph in document.paragraphs if paragraph.text == "١. حاشية صغيرة")
    assert 'w:jc w:val="start"' in body._p.pPr.xml
    assert 'w:jc w:val="start"' in note._p.pPr.xml
    assert body.runs[0].font.size.pt > note.runs[0].font.size.pt


def test_markdown_export_contains_book_text_without_internal_page_metadata(
    library: Library, tmp_path: Path,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "markdown-export.pdf", 2), name="كتاب عربي")
    for number, content in (
        (1, "## الباب الأول\n\nمتن **عريض**"),
        (2, "متن الصفحة الثانية\n\n> حاشية"),
    ):
        extracted = ExtractedPage(
            is_blank=False,
            printed_page=str(number),
            content_markdown=content.split("\n\n"),
        )
        library.apply_extraction(book["id"], number, extracted, extracted.model_dump_json())
        page = library.get_page(book["id"], number)
        library.save_page_content(
            book["id"], number, page["content_html"],
            content_markdown=page["content_markdown"], reviewed=True,
        )

    destination = tmp_path / "book.md"
    report = export_markdown(library, book["id"], 1, 2, destination)
    exported = destination.read_text(encoding="utf-8")

    assert report["format"] == "markdown"
    assert exported.startswith("# كتاب عربي")
    assert "<!-- waraq-page:" not in exported
    assert "printed-page:" not in exported
    assert "## الباب الأول" in exported
    assert "متن **عريض**" in exported
    assert "> حاشية" in exported


def test_html_export_is_a_self_contained_rtl_book_with_navigation_and_search(
    library: Library, tmp_path: Path,
) -> None:
    book = library.add_book(make_pdf(tmp_path / "html-export.pdf", 2), name="كتاب البحث")
    for number, text in ((1, "مقدمة الكتاب"), (2, "العبارة المطلوبة في الفهرس")):
        extracted = result(text)
        library.apply_extraction(book["id"], number, extracted, extracted.model_dump_json())
        page = library.get_page(book["id"], number)
        library.save_page_content(book["id"], number, page["content_html"], reviewed=True)

    destination = tmp_path / "book.html"
    report = export_html(library, book["id"], 1, 2, destination)
    exported = destination.read_text(encoding="utf-8")

    assert report["format"] == "html"
    assert '<html lang="ar" dir="rtl">' in exported
    assert 'class="book-page" data-page="1"' in exported
    assert 'class="book-page" data-page="2"' in exported
    assert 'id="previous-page"' in exported
    assert 'id="next-page"' in exported
    assert 'id="page-number"' in exported
    assert 'id="book-search"' in exported
    assert 'id="search-results"' in exported
    assert "العبارة المطلوبة في الفهرس" in exported
    assert "addEventListener(\"keydown\"" in exported
    assert "ArrowRight" in exported and "ArrowLeft" in exported
    assert "data:text" not in exported
    assert not re.search(r'<(?:script|link)[^>]+(?:src|href)=["\']https?://', exported)


def test_word_export_can_collect_only_reviewed_pages_from_an_incomplete_book(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 4))
    for number in (1, 3):
        extracted = result(f"صفحة معتمدة {number}")
        library.apply_extraction(book["id"], number, extracted, extracted.model_dump_json())
        page = library.get_page(book["id"], number)
        library.save_regions(book["id"], number, page["regions"], reviewed=True)

    monkeypatch.setattr("waraq.exporter._office_path", lambda: None)
    destination = tmp_path / "approved.docx"
    report = export_word(
        library,
        book["id"],
        1,
        4,
        destination,
    )

    assert report["source_pages"] == [1, 3]
    assert report["skipped_pages"] == [2, 4]
    assert report["expected_pages"] == 2
    with ZipFile(destination) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "صفحة معتمدة 1" in xml and "صفحة معتمدة 3" in xml


def test_word_export_can_include_generated_unreviewed_pages_and_skips_ungenerated_pages(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 4))
    for number in (1, 2):
        extracted = result(f"صفحة مولّدة {number}")
        library.apply_extraction(book["id"], number, extracted, extracted.model_dump_json())
    page = library.get_page(book["id"], 1)
    library.save_regions(book["id"], 1, page["regions"], reviewed=True)

    monkeypatch.setattr("waraq.exporter._office_path", lambda: None)
    destination = tmp_path / "generated.docx"
    report = export_word(
        library,
        book["id"],
        1,
        4,
        destination,
        include_unreviewed=True,
    )

    assert report["source_pages"] == [1, 2]
    assert report["reviewed_pages"] == 1
    assert report["unreviewed_pages"] == 1
    assert report["skipped_unconverted_pages"] == [3, 4]
    with ZipFile(destination) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "صفحة مولّدة 1" in xml and "صفحة مولّدة 2" in xml


def test_word_export_preserves_html_heading_level_as_a_word_heading_style(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
):
    book = library.add_book(make_pdf(tmp_path / "heading-source.pdf", 1))
    extracted = ExtractedPage(
        is_blank=False,
        printed_page="١",
        content_markdown=["## عنوان فرعي", "المتن"],
    )
    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())
    page = library.get_page(book["id"], 1)
    library.save_page_content(book["id"], 1, page["content_html"], reviewed=True)
    monkeypatch.setattr("waraq.exporter._office_path", lambda: None)

    destination = tmp_path / "heading.docx"
    export_word(library, book["id"], 1, 1, destination)

    with ZipFile(destination) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert 'w:pStyle w:val="Heading2"' in xml


def test_extracted_line_removes_layout_breaks_inside_one_visual_line():
    extracted = result("\n\n  بسم الله   الرحمن الرحيم \n\n")
    assert extracted.content_markdown[0] == "بسم الله   الرحمن الرحيم"


def test_bridge_reports_background_export_start_and_completion(
    library: Library, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    application = QGuiApplication.instance() or QGuiApplication([])
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 1))
    calls = []

    def fake_export(*args, **kwargs):
        calls.append(kwargs)
        return {"message": "اكتمل التصدير"}

    monkeypatch.setattr("waraq.bridge.export_word", fake_export)
    bridge = Bridge(library, MemorySecretStore())
    bridge.openBook(book["id"])
    bridge.exportBook(json.dumps({
        "book_id": book["id"], "start_page": 1, "end_page": 1,
        "destination": str(tmp_path / "book.docx"), "include_unreviewed": True,
    }))
    assert bridge.exporting
    assert bridge.toast.startswith("بدأ إنشاء ملف Word")

    deadline = time.monotonic() + 1
    while bridge.exporting and time.monotonic() < deadline:
        application.processEvents()
        time.sleep(.01)

    assert not bridge.exporting
    assert calls[0]["include_unreviewed"] is True
    assert "اكتمل التصدير" in bridge.toast
    assert str(tmp_path / "book.docx") in bridge.toast


@pytest.mark.parametrize(
    ("export_format", "attribute", "suffix"),
    [("markdown", "export_markdown", ".md"), ("html", "export_html", ".html")],
)
def test_bridge_dispatches_text_exports_and_corrects_the_destination_suffix(
    library: Library,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    export_format: str,
    attribute: str,
    suffix: str,
) -> None:
    application = QGuiApplication.instance() or QGuiApplication([])
    book = library.add_book(make_pdf(tmp_path / f"{export_format}.pdf", 1))
    calls: list[Path] = []

    def fake_export(_library, _book_id, _start, _end, destination, **_kwargs):
        calls.append(destination)
        return {"message": "اكتمل التصدير"}

    monkeypatch.setattr(f"waraq.bridge.{attribute}", fake_export)
    bridge = Bridge(library, MemorySecretStore())
    bridge.openBook(book["id"])
    bridge.exportBook(json.dumps({
        "book_id": book["id"],
        "start_page": 1,
        "end_page": 1,
        "destination": str(tmp_path / "book.docx"),
        "format": export_format,
    }))

    deadline = time.monotonic() + 1
    while bridge.exporting and time.monotonic() < deadline:
        application.processEvents()
        time.sleep(.01)

    assert calls == [tmp_path / f"book{suffix}"]


def test_page_rich_editor_persists_one_document_and_plain_line_compatibility(
    library: Library, tmp_path: Path
):
    application = QGuiApplication.instance() or QGuiApplication([])
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 1))
    extracted = result("النص القديم")
    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())
    bridge = Bridge(library, MemorySecretStore())
    bridge.openBook(book["id"])
    bridge.openPage(1)
    bridge.updatePageRichText(
        '<p><strong>عنوان مصحح</strong></p>'
        '<p><span class="ql-font-naskh ql-size-huge"><u>النص المصحح</u></span></p>'
        '<blockquote>اقتباس جديد</blockquote><ol><li data-list="ordered">بند أول</li></ol>'
        '<img src="https://example.invalid/tracker.png"><script>alert(1)</script>'
    )
    bridge.savePage()
    page = library.get_page(book["id"], 1)
    assert "<strong>عنوان مصحح</strong>" in page["content_html"]
    assert "<u>النص المصحح</u>" in page["content_html"]
    assert "<blockquote>اقتباس جديد</blockquote>" in page["content_html"]
    assert 'class="ql-font-naskh ql-size-huge"' in page["content_html"]
    assert 'data-list="ordered"' in page["content_html"]
    assert "script" not in page["content_html"] and "img" not in page["content_html"]
    assert [line["text"] for line in page["regions"][0]["lines"]] == [
        "عنوان مصحح", "النص المصحح", "اقتباس جديد", "بند أول",
    ]
    assert len(page["regions"]) == 1

    reopened = Bridge(library, MemorySecretStore())
    reopened.openBook(book["id"])
    reopened.openPage(1)
    assert reopened.currentPage["content_html"] == page["content_html"]


def test_markdown_editor_saves_markdown_as_canonical_and_derives_safe_html(
    library: Library, tmp_path: Path,
) -> None:
    QGuiApplication.instance() or QGuiApplication([])
    book = library.add_book(make_pdf(tmp_path / "markdown-editor.pdf", 1))
    extracted = result("السطر الأصلي")
    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())
    bridge = Bridge(library, MemorySecretStore())
    bridge.openBook(book["id"])
    bridge.openPage(1)

    bridge.updatePageEditorContent(
        "markdown",
        "السطر الأول\n\n## عنوان\n\n> حاشية صغيرة",
        "<script>لا تحفظ هذا</script>",
    )
    bridge.savePage()

    page = library.get_page(book["id"], 1)
    assert page["content_format"] == "markdown"
    assert page["content_markdown"] == "السطر الأول\n\n## عنوان\n\n> حاشية صغيرة"
    assert '<h2 class="ql-align-center">عنوان</h2>' in page["content_html"]
    assert '<blockquote class="ql-size-small">' in page["content_html"]
    assert "script" not in page["content_html"]


def test_bridge_reset_cancels_pending_autosave_and_reloads_the_original_text(
    library: Library, tmp_path: Path,
) -> None:
    application = QGuiApplication.instance() or QGuiApplication([])
    book = library.add_book(make_pdf(tmp_path / "bridge-reset.pdf", 1))
    extracted = result("التحويل الأصلي")
    library.apply_extraction(
        book["id"], 1, extracted, extracted.model_dump_json()
    )
    bridge = Bridge(library, MemorySecretStore())
    bridge.openBook(book["id"])
    bridge.openPage(1)
    bridge.updatePageRichText("<p>تعديل لم يحن وقت حفظه</p>")
    assert bridge._autosave.isActive()

    bridge.resetCurrentPageToExtraction()
    application.processEvents()

    assert not bridge._autosave.isActive()
    assert bridge.currentPage["content_html"] == "<p>التحويل الأصلي</p>"
    assert library.get_page(book["id"], 1)["content_html"] == (
        "<p>التحويل الأصلي</p>"
    )
    assert bridge.toast == "أُعيد نص الصفحة إلى التحويل الأصلي"


def test_approving_reads_the_latest_rich_editor_html_before_delayed_autosave(
    library: Library, tmp_path: Path
):
    application = QGuiApplication.instance() or QGuiApplication([])
    book = library.add_book(make_pdf(tmp_path / "source.pdf", 1))
    extracted = result("النص القديم")
    library.apply_extraction(book["id"], 1, extracted, extracted.model_dump_json())
    bridge = Bridge(library, MemorySecretStore())
    bridge.openBook(book["id"])
    bridge.openPage(1)

    latest_html = (
        '<h1 class="ql-align-right"><span class="ql-font-naskh" '
        'style="color: #6256d9; background-color: #fff1a8">عنوان منسق</span></h1>'
        '<ol><li data-list="ordered"><strong><u>بند محفوظ</u></strong></li></ol>'
    )
    bridge.commitPageRichText(book["id"], 1, latest_html, True)

    reopened = Bridge(library, MemorySecretStore())
    reopened.openBook(book["id"])
    reopened.openPage(1)
    saved = reopened.currentPage["content_html"]
    assert reopened.currentPage["reviewed"] == 1
    assert 'class="ql-align-right"' in saved
    assert 'class="ql-font-naskh"' in saved
    assert "background-color: #fff1a8" in saved
    assert "<strong><u>بند محفوظ</u></strong>" in saved


def test_toast_is_scheduled_to_close_after_five_seconds(library: Library):
    application = QGuiApplication.instance() or QGuiApplication([])
    bridge = Bridge(library, MemorySecretStore())
    bridge._notify("اكتمل الاختبار")

    assert bridge.toast == "اكتمل الاختبار"
    assert bridge._toast_timer.interval() == 5_000
    assert bridge._toast_timer.isActive()

    bridge.clearToast()
    assert not bridge.toast
    assert not bridge._toast_timer.isActive()

    bridge._toast_timer.setInterval(10)
    bridge._notify("تنبيه قصير للاختبار")
    deadline = time.monotonic() + .25
    while bridge.toast and time.monotonic() < deadline:
        application.processEvents()
        time.sleep(.005)
    assert not bridge.toast


def test_task_can_return_to_the_page_it_was_created_for(library: Library, tmp_path: Path):
    first = library.add_book(make_pdf(tmp_path / "first.pdf", 2), "الأول")
    second = library.add_book(make_pdf(tmp_path / "second.pdf", 4), "الثاني")
    key = library.add_key("المفتاح", "secret-ref")
    task_id = library.create_task(
        book_id=second["id"], start_page=3, end_page=3,
        key_id=key, model="gemini-test", prompt_id="default-printed",
        mode="printed", dpi=300, overwrite=False,
    )
    bridge = Bridge(library, MemorySecretStore())
    bridge.openBook(first["id"])
    bridge._task = library.task(task_id)

    bridge.openTaskPage()

    assert bridge.currentBook["id"] == second["id"]
    assert bridge.currentPage["number"] == 3
    assert bridge.screen == "review"


def test_extraction_schema_rejects_text_on_a_blank_page():
    with pytest.raises(ValueError):
        ExtractedPage(is_blank=True, printed_page="", content_markdown=["نص"])
