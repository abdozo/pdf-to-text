"""Regression cases from reviewing real summary selections."""
import json
from types import SimpleNamespace

import pytest

from tests_desktop.test_summaries import setup, task, saved, document
from waraq.tasks import ConversionRunner
from waraq.database import Library
from waraq.gemini_client import GeminiFailure
from waraq.summaries import (
    MAX_REVIEW_BYTES, QUADRANTS, SUMMARY_PROMPT, SUMMARY_REVIEW_PROMPT,
    SummaryDocument, json_bytes, review_plan, export_summaries, validate_review,
)
from tests_desktop.test_summaries import fake_client


def test_large_selection_starts_review_instead_of_rejecting_it(setup, monkeypatch):
    first = saved(setup, [1, 2])
    second = saved(setup, [3, 4])
    library = setup[0]
    with library.transaction() as db:
        for summary_id, pages in [(first, [1, 2]), (second, [3, 4])]:
            value = document(pages)
            value.sheets[0].understanding[0].text = "تفصيل " * 6000
            db.execute("UPDATE summaries SET document=? WHERE id=?", (value.model_dump_json(), summary_id))
    review_id = task(setup, summary_ids=[first, second])
    assert setup[0].task(review_id)["state"] == "queued"
    units = setup[0].review_units(review_id)
    assert len(units) > 2
    assert all(json_bytes(unit["target"]) < MAX_REVIEW_BYTES // 4 for unit in units)
    for summary_id in (first, second):
        text = "".join(item["text"] for unit in units if unit["source_id"] == summary_id
                       for sheet in unit["target"]["sheets"] for item in sheet["understanding"])
        assert text == ("تفصيل " * 6000).strip()


def test_review_keeps_selected_summaries_separate(setup, monkeypatch):
    library, book, _key, secrets = setup
    ids = [saved(setup, [1, 2]), saved(setup, [3, 4])]
    review_id = task(setup, summary_ids=ids)

    class Client:
        def __init__(self, _key): pass
        def summarize(self, _images, **kwargs):
            # The old path explicitly asks for one document covering the selection.
            result = document(kwargs["source_pages"])
            return SimpleNamespace(document=result, raw=result.model_dump_json(), usage={})
        def review_summary(self, *, target, **kwargs):
            result = document(target["source_pages"], "نتيجة مراجعة مستقلة")
            return SimpleNamespace(document=result, raw=result.model_dump_json(), usage={})

    monkeypatch.setattr("waraq.tasks.GeminiClient", Client)
    runner = ConversionRunner(library, secrets)
    runner.start(review_id); runner.join(5)
    assert library.task(review_id)["state"] == "completed"
    outputs = [r for r in library.list_summaries(book["id"]) if r["task_id"] == review_id]
    assert len(outputs) == 2
    assert [r["source_pages"] for r in outputs] == [[1, 2], [3, 4]]


def test_review_pause_restart_and_failure_preserve_context_and_progress(setup, monkeypatch):
    library, book, _key, secrets = setup
    ids = [saved(setup, [1]), saved(setup, [2]), saved(setup, [3])]
    review_id = task(setup, summary_ids=ids)
    calls = []
    fail = True
    class Client:
        def __init__(self, _key): pass
        def review_summary(self, **kwargs):
            calls.append(kwargs)
            if fail and kwargs["target"]["source_pages"] == [2]:
                raise GeminiFailure("temporary error")
            result = document(kwargs["target"]["source_pages"], "مراجع")
            return SimpleNamespace(document=result, raw=result.model_dump_json(), usage={})
    monkeypatch.setattr("waraq.tasks.GeminiClient", Client)
    def pause_after_save(event):
        if event.get("summary_id"):
            library.request_task_action(review_id, "pause")
    runner = ConversionRunner(library, secrets, pause_after_save)
    runner.start(review_id); runner.join(5)
    assert library.task(review_id)["state"] == "paused"
    assert library.task(review_id)["completed"] == 1
    restored = Library(library.root)
    runner = ConversionRunner(restored, secrets)
    runner.start(review_id); runner.join(5)
    assert library.task(review_id)["state"] == "failed"
    assert library.task(review_id)["completed"] == 1
    assert len(calls) == 2  # No request for unit 3 before resolving unit 2.
    fail = False
    runner.start(review_id); runner.join(5)
    assert library.task(review_id)["state"] == "completed"
    assert library.task(review_id)["completed"] == 3
    assert [call["target"]["source_pages"] for call in calls] == [[1], [2], [2], [3]]
    assert calls[2]["previous"]["sheets"][0]["title"] == "مراجع"
    outputs = [record for record in library.list_summaries(book["id"]) if record["task_id"] == review_id]
    assert len(outputs) == 3


def test_review_cancel_ignores_late_result(setup, monkeypatch):
    import threading
    library, book, _key, secrets = setup
    review_id = task(setup, summary_ids=[saved(setup, [1]), saved(setup, [2])])
    runner = ConversionRunner(library, secrets)
    cancelled = threading.Event()
    class Client:
        def __init__(self, _key): pass
        def review_summary(self, **kwargs):
            runner.request_cancel(review_id)
            cancelled.set()
            result = document([1])
            return SimpleNamespace(document=result, raw=result.model_dump_json(), usage={})
    monkeypatch.setattr("waraq.tasks.GeminiClient", Client)
    runner.start(review_id)
    # request_cancel detaches the thread; wait until it has also closed its request.
    assert cancelled.wait(3)
    assert library.task(review_id)["state"] == "cancelled"
    assert not [record for record in library.list_summaries(book["id"]) if record["task_id"] == review_id]
    assert library.request_rows()[0]["state"] == "cancelled"


def test_review_citations_may_link_neighbors_but_not_copy_their_whole_range():
    target, following = document([1, 2]).model_dump(), document([3, 4]).model_dump()
    result = document([1, 2, 3])
    result.sheets[0].critique[0].pdf_pages = [3]
    validate_review(result, target, None, following)
    result.source_pages.append(4)
    with pytest.raises(ValueError, match="نطاقات السياق"):
        validate_review(result, target, None, following)


def test_review_sdk_and_direct_writing_instructions():
    target, following = document([1]).model_dump(), document([2]).model_dump()
    result = document([1])
    client, calls = fake_client(result.model_dump_json())
    response = client.review_summary(model="gemini-3.5-flash-lite", prompt=SUMMARY_REVIEW_PROMPT,
                                     target=target, previous=None, following=following)
    assert response.document.source_pages == [1]
    sent = json.loads(calls[0]["contents"][1])
    assert sent == {"target": target, "previous": None, "following": following}
    assert "اكتب الفكرة والنتيجة مباشرة" in SUMMARY_PROMPT
    assert "لا تبدأ الفقرات" in SUMMARY_PROMPT
    assert "عدد أوراق الناتج ليس ثابتًا" in SUMMARY_REVIEW_PROMPT


def test_summary_prompt_prefers_information_over_filling_quadrants():
    assert "الربع الأول وحده استيعابي" in SUMMARY_PROMPT
    assert "يجوز أن تكون قائمة الربع فارغة" in SUMMARY_PROMPT
    assert "لا تكتب أن البحث لم يجد" in SUMMARY_PROMPT
    assert "لا تمدح المؤلف أو الكتاب" in SUMMARY_PROMPT
    assert "المعلومات المشهورة أو المسلّمات" in SUMMARY_PROMPT
    assert "تصورًا أمينًا لمحتوى المبحث كله" in SUMMARY_PROMPT
    assert "إذا لم يوجد إشكال محدد فاترك critique فارغة" in SUMMARY_PROMPT
    assert "إذا لم يوجد تطبيق نافع فاترك applications فارغة" in SUMMARY_PROMPT


@pytest.mark.parametrize("suffix", ["md", "html", "docx"])
def test_selected_export_contains_all_selected_sheets(tmp_path, suffix):
    docs = [document([2], "العنوان الأول"), document([4], "العنوان الثاني")]
    docs[1].sheets.append(document([4], "الورقة الثالثة").sheets[0])
    path = tmp_path / ("selected." + suffix)
    export_summaries([{"document": doc.model_dump()} for doc in docs], path)
    if suffix == "docx":
        from docx import Document
        result = Document(path)
        assert len(result.tables) == 3
        text = " ".join(paragraph.text for paragraph in result.paragraphs)
    else:
        text = path.read_text()
    assert all(title in text for title in ["العنوان الأول", "العنوان الثاني", "الورقة الثالثة"])


def test_configurable_batches_and_flexible_sheet_count(setup, monkeypatch):
    library, book, _key, secrets = setup
    ids = [saved(setup, [p]) for p in [1, 2, 3, 4, 5]]
    review_id = task(setup, summary_ids=ids, review_batch_size=3)
    assert library.task(review_id)['review_batch_size'] == 3
    assert [len(u['source_ids']) for u in library.review_units(review_id)] == [3, 2]
    calls = []
    class Client:
        def __init__(self, _key): pass
        def review_summary(self, **kwargs):
            calls.append(kwargs)
            pages = kwargs['target']['source_pages']
            result = document(pages)
            # Three input sheets become two; two become three.
            result.sheets = [document([p]).sheets[0] for p in ([1, 3] if pages[0] == 1 else [4, 5, 5])]
            return SimpleNamespace(document=result, raw=result.model_dump_json(), usage={})
    monkeypatch.setattr('waraq.tasks.GeminiClient', Client)
    runner = ConversionRunner(library, secrets)
    runner.start(review_id); runner.join(5)
    assert library.task(review_id)['state'] == 'completed'
    units = library.review_units(review_id)
    assert [len(u['result']['sheets']) for u in units] == [2, 3]
    assert calls[1]['previous']['sheets'] == [units[0]['result']['sheets'][-1]]
    outputs = [r for r in library.list_summaries(book['id']) if r['task_id'] == review_id]
    assert [r['input_ids'] for r in outputs] == [ids[:3], ids[3:]]


@pytest.mark.parametrize('review', [False, True])
@pytest.mark.parametrize('repair_succeeds', [False, True])
def test_style_repair_before_saving_initial_and_reviewed_summaries(setup, monkeypatch, review, repair_succeeds):
    library, book, _key, secrets = setup
    ids = [saved(setup, [1, 2])] if review else None
    task_id = task(setup, summary_ids=ids)
    repairs = []
    class Client:
        def __init__(self, _key): pass
        def summarize(self, _images, **kwargs):
            return self.response(kwargs['source_pages'])
        def review_summary(self, **kwargs):
            return self.response(kwargs['target']['source_pages'])
        def response(self, pages):
            doc = document(pages)
            doc.sheets[0].understanding[0].text = 'عرّف المصنف علم الحديث بأنه معرفة القواعد.'
            return SimpleNamespace(document=doc, raw=doc.model_dump_json(), usage={})
        def polish_summary(self, doc, **kwargs):
            repairs.append(doc.model_dump())
            doc = doc.model_copy(deep=True)
            if repair_succeeds:
                doc.sheets[0].understanding[0].text = 'علم الحديث: معرفة القواعد.'
            return SimpleNamespace(document=doc, raw=doc.model_dump_json(), usage={})
    monkeypatch.setattr('waraq.tasks.GeminiClient', Client)
    runner = ConversionRunner(library, secrets)
    runner.start(task_id); runner.join(10)
    outputs = [r for r in library.list_summaries(book['id']) if r['task_id'] == task_id]
    if repair_succeeds:
        assert library.task(task_id)['state'] == 'completed'
        assert outputs and repairs
        assert all(r['document']['sheets'][0]['understanding'][0]['text'] == 'علم الحديث: معرفة القواعد.' for r in outputs)
    else:
        assert outputs
        assert library.task(task_id)["state"] == "completed"
        assert all(r["style_warning"] for r in outputs)
        assert len(repairs) == (2 if review else 6)
    assert len(library.request_rows()) == len(repairs) + (1 if review else 3)


def test_style_edits_preserve_quotes_numbers_and_references():
    from waraq.summaries import summary_text_slots, SummaryTextEdits, apply_style_edits, author_framing
    doc = document([1])
    doc.sheets[0].understanding[0].text = 'قول المؤلف حرفيًا: «العلم نافع» وله 3 أقسام.'
    assert author_framing(doc)
    edits = {'edits': [{'id': i, 'text': getattr(obj, key)} for i, (obj, key) in enumerate(summary_text_slots(doc))]}
    edits['edits'][2]['text'] = '«العلم نافع» وله 3 أقسام.'
    revised = apply_style_edits(doc, SummaryTextEdits.model_validate(edits))
    assert not author_framing(revised)
    assert revised.sheets[0].understanding[0].pdf_pages == [1]
    for text in ['«العلم ضار» وله 3 أقسام.', '«العلم نافع» وله 4 أقسام.']:
        edits['edits'][2]['text'] = text
        with pytest.raises(ValueError):
            apply_style_edits(doc, SummaryTextEdits.model_validate(edits))
    doc.sheets[0].understanding[0].text = '«قال المؤلف: العلم نافع»'
    assert not author_framing(doc)


def test_polish_sdk_edits_only_text_slots():
    from waraq.summaries import summary_text_slots
    doc = document([1])
    doc.sheets[0].understanding[0].text = 'عرّف المصنف العلم بأنه معرفة.'
    edits = {'edits': [{'id': i, 'text': getattr(obj, key)} for i, (obj, key) in enumerate(summary_text_slots(doc))]}
    edits['edits'][2]['text'] = 'العلم: معرفة.'
    client, calls = fake_client(json.dumps(edits))
    result = client.polish_summary(doc, model='gemini-3.5-flash-lite')
    assert result.document.sheets[0].understanding[0].text == 'العلم: معرفة.'
    assert result.document.source_pages == [1]
    assert 'pdf_pages' not in calls[0]['contents'][1]
    assert doc.sheets[0].understanding[0].text == 'عرّف المصنف العلم بأنه معرفة.'


def test_delete_review_preserves_originals_other_reviews_and_usage(setup):
    library, book, _key, _secrets = setup
    original = saved(setup, [1, 2])
    reviews = []
    for _ in range(2):
        review_id = task(setup, summary_ids=[original])
        library.prepare_task_run(review_id)
        library.start_review_unit(review_id, 0)
        library.save_review_unit(review_id, 0, document([1, 2]))
        request_id = library.start_request(library.task(review_id), [1, 2])
        library.finish_request(request_id, 'success', usage={})
        library.update_task(review_id, state='completed')
        reviews.append(review_id)
    before = library.request_rows()
    library.delete_summary_review(book['id'], reviews[0])
    remaining = library.list_summaries(book['id'])
    assert original in [r['id'] for r in remaining]
    assert [r['task_id'] for r in remaining if r['input_ids']] == reviews[1:]
    assert library.request_rows() == before
    assert not library.review_units(reviews[0])
    with pytest.raises(RuntimeError):
        library.prepare_task_run(reviews[0])
    with pytest.raises(ValueError):
        library.delete_summary_review(book['id'], remaining[0]['task_id'])


def test_delete_review_rejects_running_and_unfinished_dependents(setup):
    library, book, _key, _secrets = setup
    original = saved(setup, [1])
    review_id = task(setup, summary_ids=[original])
    library.prepare_task_run(review_id)
    library.start_review_unit(review_id, 0)
    result = library.save_review_unit(review_id, 0, document([1]))
    with pytest.raises(ValueError, match='علّق'):
        library.delete_summary_review(book['id'], review_id)
    library.update_task(review_id, state='completed')
    dependent = task(setup, summary_ids=[result])
    with pytest.raises(ValueError, match='تعتمد'):
        library.delete_summary_review(book['id'], review_id)
    library.request_task_action(dependent, 'cancel')
    library.delete_summary_review(book['id'], review_id)


@pytest.mark.parametrize('reported', [[1, 2], [1, 1, 2, 3]])
def test_review_rebuilds_source_index_from_valid_neighbor_citations(reported):
    target = document([1, 2]).model_dump()
    following = document([3, 4]).model_dump()
    result = document(reported)
    result.sheets[0].critique[0].pdf_pages = [3]
    before = result.sheets[0].model_dump()
    validate_review(result, target, None, following)
    assert result.source_pages == [1, 2, 3]
    assert result.sheets[0].model_dump() == before


def test_review_source_index_repair_still_rejects_invented_citations():
    result = document([1, 2])
    result.sheets[0].critique[0].pdf_pages = [99]
    with pytest.raises(ValueError, match='إحالة'):
        validate_review(result, document([1, 2]).model_dump(), None, document([3]).model_dump())
    assert result.source_pages == [1, 2]


def test_review_sdk_accepts_verified_neighbor_missing_from_source_index():
    result = document([1, 2])
    result.sheets[0].critique[0].pdf_pages = [3]
    client, _calls = fake_client(result.model_dump_json())
    response = client.review_summary(model='gemini-3.5-flash-lite', prompt=SUMMARY_REVIEW_PROMPT,
                                     target=document([1, 2]).model_dump(), following=document([3]).model_dump())
    assert response.document.source_pages == [1, 2, 3]


@pytest.mark.parametrize('review', [False, True])
def test_delete_selected_preserves_unselected_results(setup, review):
    library, book, _key, _secrets = setup
    originals = [saved(setup, [1]), saved(setup, [2])]
    for row in library.list_summaries(book['id']):
        library.update_task(row['task_id'], state='completed')
    chosen = originals
    if review:
        review_id = task(setup, summary_ids=originals)
        library.prepare_task_run(review_id)
        chosen = []
        for i in range(2):
            library.start_review_unit(review_id, i)
            chosen.append(library.save_review_unit(review_id, i, document([i + 1])))
        library.update_task(review_id, state='completed')
    library.delete_summaries(book['id'], [chosen[0]])
    remaining = {r['id'] for r in library.list_summaries(book['id'])}
    assert chosen[0] not in remaining
    assert chosen[1] in remaining
    if review:
        assert set(originals) <= remaining
    assert library.page_path(book['id']).exists()
    library.delete_summaries(book['id'], [chosen[1]])
    assert {r['id'] for r in library.list_summaries(book['id'])} == (set(originals) if review else set())


def test_delete_selected_is_atomic_and_cancels_paused_producer(setup):
    library, book, _key, _secrets = setup
    original = saved(setup, [1])
    producer = library.list_summaries(book['id'])[0]['task_id']
    with pytest.raises(ValueError):
        library.delete_summaries(book['id'], [original, 'missing'])
    assert len(library.list_summaries(book['id'])) == 1
    with pytest.raises(ValueError, match='علّق'):
        library.delete_summaries(book['id'], [original])
    library.update_task(producer, state='paused')
    library.delete_summaries(book['id'], [original])
    assert library.task(producer)['state'] == 'cancelled'


def test_attribution_required_in_initial_review_and_editorial_prompts():
    from waraq.summaries import ATTRIBUTION_RULES, SUMMARY_STYLE_PROMPT, SUMMARY_SYSTEM
    for prompt in (SUMMARY_PROMPT, SUMMARY_REVIEW_PROMPT, SUMMARY_STYLE_PROMPT, SUMMARY_SYSTEM):
        assert ATTRIBUTION_RULES in prompt
        assert 'قول الإمام أحمد بن حنبل' in prompt
        assert 'حديث نبوي' in prompt
        assert 'لا تستعدها من الذاكرة' in prompt
    assert 'بين علامتي تنصيص دون مقدمة' not in SUMMARY_PROMPT
    assert 'أسماء أصحاب الأقوال فقط عندما' not in SUMMARY_PROMPT


def test_delete_after_cancellation_handles_older_stopped_reviews(setup):
    library, book, _key, _secrets = setup
    original = saved(setup, [1])
    producer = library.list_summaries(book['id'])[0]['task_id']
    library.update_task(producer, state='completed')
    older = task(setup, summary_ids=[original])
    library.update_task(older, state='failed')
    latest = task(setup, summary_ids=[original])
    library.request_task_action(latest, 'cancel')
    library.delete_summaries(book['id'], [original])
    assert not library.list_summaries(book['id'])
    assert library.task(older)['state'] == 'cancelled'
    assert library.task(latest)['state'] == 'cancelled'


def test_running_dependent_blocks_delete_without_cancelling_others(setup):
    library, book, _key, _secrets = setup
    original = saved(setup, [1])
    producer = library.list_summaries(book['id'])[0]['task_id']
    library.update_task(producer, state='completed')
    stopped = task(setup, summary_ids=[original])
    library.update_task(stopped, state='failed')
    running = task(setup, summary_ids=[original])
    library.prepare_task_run(running)
    with pytest.raises(ValueError, match='تعمل الآن'):
        library.delete_summaries(book['id'], [original])
    assert library.task(stopped)['state'] == 'failed'
    assert library.task(running)['state'] == 'running'
    assert library.list_summaries(book['id'])[0]['id'] == original
