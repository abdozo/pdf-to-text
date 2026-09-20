from __future__ import annotations

import html
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
import threading
import unicodedata
import uuid
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from .database import Library, utc_now
from .richtext import rich_html_blocks, sanitize_rich_html


EXPORT_LOCK = threading.Lock()
ALLOWED_TAG = re.compile(r"(<sup>[^<>]*</sup>)")


def _office_path() -> str | None:
    configured = os.environ.get("WARRAQ_SOFFICE")
    program_files = os.environ.get("PROGRAMFILES")
    program_files_x86 = os.environ.get("PROGRAMFILES(X86)")
    local_app_data = os.environ.get("LOCALAPPDATA")
    candidates = (
        configured,
        shutil.which("soffice"),
        str(Path(program_files) / "LibreOffice" / "program" / "soffice.exe") if program_files else None,
        str(Path(program_files_x86) / "LibreOffice" / "program" / "soffice.exe") if program_files_x86 else None,
        str(Path(local_app_data) / "Programs" / "LibreOffice" / "program" / "soffice.exe") if local_app_data else None,
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        str(Path(__file__).resolve().parent / "runtime" / "LibreOffice.app" / "Contents" / "MacOS" / "soffice"),
    )
    return next((str(path) for path in candidates if path and Path(path).is_file()), None)


def _set_cell_margins(cell: Any, top: int = 60, start: int = 80, bottom: int = 60, end: int = 80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def _set_paragraph_rtl(paragraph: Any) -> None:
    """Write portable RTL paragraph properties in the OOXML schema order."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
    else:
        ppr.remove(bidi)
    bidi.set(qn("w:val"), "1")
    ppr.insert_element_before(
        bidi,
        "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap",
        "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
        "w:outlineLvl", "w:divId", "w:cnfStyle",
    )
    ppr.find(qn("w:jc")).set(qn("w:val"), "start")


def _uses_rtl(text: str) -> bool:
    for character in text:
        direction = unicodedata.bidirectional(character)
        if direction in {"R", "AL"}:
            return True
        if direction == "L":
            return False
    return True


def _set_paragraph_ltr(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.insert_element_before(bidi, "w:jc")
    bidi.set(qn("w:val"), "0")


def _set_paragraph_alignment(paragraph: Any, alignment: str, *, rtl: bool = True) -> None:
    """Use logical alignment so RTL stays portable across Word renderers."""
    value = {
        "right": "start" if rtl else "end",
        "left": "end" if rtl else "start",
        "center": "center",
        "justify": "both",
    }.get(alignment)
    if value is None:
        return
    ppr = paragraph._p.get_or_add_pPr()
    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        ppr.append(jc)
    jc.set(qn("w:val"), value)


def _set_default_paragraph_rtl(document: Any) -> None:
    """Make right alignment the Word fallback when a paragraph has no override."""
    normal = document.styles["Normal"]
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = normal._element.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
    else:
        ppr.remove(bidi)
    bidi.set(qn("w:val"), "1")
    ppr.insert_element_before(
        bidi,
        "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap",
        "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
        "w:outlineLvl", "w:divId", "w:cnfStyle",
    )
    ppr.find(qn("w:jc")).set(qn("w:val"), "start")


def _set_run_rtl(run: Any) -> None:
    rpr = run._r.get_or_add_rPr()
    for tag in ("rtl", "cs"):
        element = rpr.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            rpr.append(element)
        element.set(qn("w:val"), "1")
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is not None:
        fonts.set(qn("w:cs"), run.font.name or "Arial")
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:bidi"), "ar-SA")


def _rtl(paragraph: Any, text: str, size: float, *, handwritten: bool = False, bold: bool = False, text_direction: str = "auto") -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = 1.15
    rtl = text_direction == "rtl" or (text_direction == "auto" and _uses_rtl(text))
    (_set_paragraph_rtl if rtl else _set_paragraph_ltr)(paragraph)
    for part in ALLOWED_TAG.split(text):
        superscript = part.startswith("<sup>") and part.endswith("</sup>")
        value = part[5:-6] if superscript else re.sub(r"<[^>]+>", "", part)
        run = paragraph.add_run(value)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.superscript = superscript
        run.font.italic = handwritten
        run.font.bold = bold
        if rtl:
            _set_run_rtl(run)


def _rtl_rich(paragraph: Any, block: dict[str, Any], size: float, scale: float, *, text_direction: str = "auto") -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = 1.15
    if block.get("quote"):
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.right_indent = Pt(18)
    heading_level = block.get("heading_level")
    if heading_level:
        paragraph.style = f"Heading {heading_level}"
    list_kind = next((run.get("list") for run in block.get("runs", []) if run.get("list")), None)
    if list_kind:
        paragraph.style = "List Number" if list_kind == "ordered" else "List Bullet"
    rtl = text_direction == "rtl" or (
        text_direction == "auto"
        and _uses_rtl("".join(item.get("text", "") for item in block.get("runs", [])))
    )
    (_set_paragraph_rtl if rtl else _set_paragraph_ltr)(paragraph)
    if block.get("separator"):
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in {
            "val": "single", "sz": "4", "space": "1", "color": "777777",
        }.items():
            bottom.set(qn(f"w:{key}"), value)
        border.append(bottom)
        paragraph._p.get_or_add_pPr().append(border)
        return
    if block.get("align"):
        _set_paragraph_alignment(paragraph, block["align"], rtl=rtl)
    for item in block.get("runs", []):
        run = paragraph.add_run(item.get("text", ""))
        run.font.name = item.get("font_family") or "Arial"
        requested_size = float(item.get("font_size") or size)
        if block.get("quote"):
            requested_size = min(requested_size, size * 0.82)
        run.font.size = Pt(requested_size * scale)
        run.font.bold = bool(item.get("bold"))
        run.font.italic = bool(item.get("italic"))
        run.font.underline = bool(item.get("underline"))
        run.font.superscript = bool(item.get("superscript"))
        run.font.subscript = bool(item.get("subscript"))
        if rtl:
            _set_run_rtl(run)


def _page_lines(page: dict[str, Any]) -> list[dict[str, Any]]:
    lines: list[dict[str, Any]] = []
    for region in sorted(page["regions"], key=lambda item: item["reading_order"]):
        for line in sorted(region["lines"], key=lambda item: item["reading_order"]):
            lines.append({**line, "kind": region["kind"], "source": region["source"], "column": region["column_index"]})
    return lines


def _export_selection(
    library: Library,
    book_id: str,
    start_page: int,
    end_page: int,
    include_unreviewed: bool,
) -> tuple[dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    book = library.get_book(book_id)
    if start_page < 1 or end_page > book["page_count"] or start_page > end_page:
        raise ValueError("نطاق التصدير غير صالح")
    requested_pages = [
        library.get_page(book_id, number)
        for number in range(start_page, end_page + 1)
    ]
    generated_pages = [page for page in requested_pages if page["state"] == "done"]
    pages = (
        generated_pages
        if include_unreviewed
        else [page for page in generated_pages if page["reviewed"]]
    )
    skipped_pages = [page["number"] for page in requested_pages if page not in pages]
    skipped_unconverted = [
        page["number"] for page in requested_pages if page["state"] != "done"
    ]
    skipped_unreviewed = [
        page["number"]
        for page in generated_pages
        if not page["reviewed"] and not include_unreviewed
    ]
    if not pages:
        if generated_pages:
            raise ValueError(
                "لا توجد صفحات معتمدة في النطاق المحدد. اختر تضمين الصفحات "
                "المولّدة غير المعتمدة للتصدير."
            )
        raise ValueError("لا توجد صفحات مولّدة في النطاق المحدد")
    return book, pages, {
        "generated_pages": len(generated_pages),
        "skipped_pages": skipped_pages,
        "skipped_unreviewed_pages": skipped_unreviewed,
        "skipped_unconverted_pages": skipped_unconverted,
    }


def _export_report(
    library: Library,
    book: dict[str, Any],
    pages: list[dict[str, Any]],
    selection: dict[str, Any],
    start_page: int,
    end_page: int,
    destination: Path,
    export_format: str,
    *,
    verified: bool,
    actual_pages: int | None = None,
    font_scale: float | None = None,
) -> dict[str, Any]:
    exported_unreviewed = sum(not bool(page["reviewed"]) for page in pages)
    exported_reviewed = len(pages) - exported_unreviewed
    message_parts = [f"أُنشئ الملف من {len(pages)} صفحة مولّدة"]
    if exported_unreviewed:
        message_parts.append(f"منها {exported_unreviewed} غير معتمدة")
    if selection["skipped_unreviewed_pages"]:
        message_parts.append(
            f"تُركت {len(selection['skipped_unreviewed_pages'])} صفحة مولّدة غير معتمدة"
        )
    if selection["skipped_unconverted_pages"]:
        message_parts.append(
            f"لم تُضف {len(selection['skipped_unconverted_pages'])} صفحة غير مولّدة"
        )
    message = ". ".join(message_parts) + "."
    if export_format == "word" and not verified:
        message += " لم يكتمل فحص مطابقة الصفحات؛ راجع الملف في Word قبل اعتماده."
    report = {
        "book": book["name"],
        "format": export_format,
        "start_page": start_page,
        "end_page": end_page,
        "expected_pages": len(pages),
        "actual_pages": actual_pages,
        "verified": verified,
        "reviewed_pages": exported_reviewed,
        "unreviewed_pages": exported_unreviewed,
        "generated_pages": selection["generated_pages"],
        "source_pages": [page["number"] for page in pages],
        "skipped_pages": selection["skipped_pages"],
        "skipped_unreviewed_pages": selection["skipped_unreviewed_pages"],
        "skipped_unconverted_pages": selection["skipped_unconverted_pages"],
        "complete_book": (
            start_page == 1
            and end_page == book["page_count"]
            and not selection["skipped_pages"]
        ),
        "message": message,
    }
    if font_scale is not None:
        report["font_scale"] = font_scale
    with library.transaction() as db:
        db.execute(
            "INSERT INTO exports VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                uuid.uuid4().hex,
                book["id"],
                start_page,
                end_page,
                str(destination),
                json.dumps(report, ensure_ascii=False),
                utc_now(),
            ),
        )
    return report


def build_docx(pages: list[dict[str, Any]], destination: Path, scale: float = 1.0) -> None:
    doc = Document()
    _set_default_paragraph_rtl(doc)
    doc.core_properties.title = "نص الكتاب المطابق لصفحات الأصل"
    doc.core_properties.description = "نص عربي قابل للتحرير. يحتاج إلى مراجعة بشرية مقابل الأصل."
    for index, page in enumerate(pages):
        section = doc.sections[0] if index == 0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.page_width = Pt(page["width"])
        section.page_height = Pt(page["height"])
        margin = min(30, page["width"] * 0.055, page["height"] * 0.055)
        section.top_margin = section.bottom_margin = Pt(margin)
        section.left_margin = section.right_margin = Pt(margin)
        section.header_distance = section.footer_distance = Pt(0)
        section.footer.is_linked_to_previous = False
        if page.get("printed_page"):
            _rtl(section.footer.paragraphs[0], page["printed_page"], 8.5)

        lines = _page_lines(page)
        text_direction = page.get("text_direction", "auto")
        rich_blocks = rich_html_blocks(page.get("content_html", "")) if page.get("content_html") else []
        page_height = page["height"] - 2 * margin - 18
        page_width = page["width"] - 2 * margin
        estimated_texts = (
            ["".join(run["text"] for run in block["runs"]) for block in rich_blocks]
            if rich_blocks else [line["text"] for line in lines]
        )
        estimated = sum(max(1, math.ceil(len(text) * 5.7 / max(page_width, 1))) for text in estimated_texts)
        size = max(6.5, min(13.5, page_height / max(estimated * 1.35, 1))) * scale
        body = [line for line in lines if line["kind"] not in {"footnote", "page_number", "margin_note"}]
        notes = [line for line in lines if line["kind"] in {"footnote", "margin_note"}]
        columns = max((line["column"] for line in body), default=0) + 1
        if rich_blocks:
            for block in rich_blocks:
                _rtl_rich(doc.add_paragraph(), block, size / scale, scale, text_direction=text_direction)
        elif not lines:
            _rtl(doc.add_paragraph(), "", 1)
        elif columns <= 1:
            for line in body:
                _rtl(doc.add_paragraph(), line["text"], size, handwritten=line["source"] == "handwritten", bold=line["kind"] == "heading", text_direction=text_direction)
        else:
            table = doc.add_table(rows=1, cols=columns)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            # Right-to-left visual order: logical column zero sits on the right.
            for logical in range(columns):
                cell = table.cell(0, columns - logical - 1)
                cell.width = Pt(page_width / columns)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                _set_cell_margins(cell)
                current = [line for line in body if line["column"] == logical]
                if not current:
                    current = [{"text": "", "source": "printed", "kind": "body"}]
                for line_index, line in enumerate(current):
                    paragraph = cell.paragraphs[0] if line_index == 0 else cell.add_paragraph()
                    _rtl(paragraph, line["text"], size, handwritten=line["source"] == "handwritten", bold=line["kind"] == "heading", text_direction=text_direction)
            tbl_pr = table._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                element = OxmlElement(f"w:{edge}"); element.set(qn("w:val"), "nil"); borders.append(element)
            tbl_pr.append(borders)
        if notes:
            first = True
            for line in notes:
                paragraph = doc.add_paragraph()
                if first:
                    paragraph.paragraph_format.space_before = Pt(7)
                    border = OxmlElement("w:pBdr"); top = OxmlElement("w:top")
                    for key, value in {"val": "single", "sz": "4", "space": "4", "color": "777777"}.items():
                        top.set(qn(f"w:{key}"), value)
                    border.append(top); paragraph._p.get_or_add_pPr().append(border); first = False
                _rtl(paragraph, line["text"], size * 0.82, handwritten=line["source"] == "handwritten", text_direction=text_direction)
    doc.save(destination)


def _pdf_page_count(path: Path) -> int:
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(path) as pdf:
        return len(pdf)


def export_word(
    library: Library,
    book_id: str,
    start_page: int,
    end_page: int,
    destination: Path,
    *,
    include_unreviewed: bool = False,
) -> dict[str, Any]:
    book, pages, selection = _export_selection(
        library, book_id, start_page, end_page, include_unreviewed
    )

    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    office = _office_path()
    actual_pages = None
    verified = False
    used_scale = 1.0
    with EXPORT_LOCK:
        if not office:
            build_docx(pages, destination)
        else:
            with tempfile.TemporaryDirectory(prefix="waraq-word-") as temporary:
                temp = Path(temporary)
                candidate = temp / destination.name
                profile = (temp / "profile").as_uri()
                rendered = temp / f"{candidate.stem}.pdf"
                for scale in (1.0, 0.92, 0.84, 0.76, 0.68, 0.60):
                    build_docx(pages, candidate, scale)
                    used_scale = scale
                    process = subprocess.run([office, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf", "--outdir", str(temp), str(candidate)], capture_output=True, timeout=300)
                    if process.returncode != 0 or not rendered.exists():
                        break
                    actual_pages = _pdf_page_count(rendered)
                    if actual_pages == len(pages):
                        verified = True
                        shutil.copy2(candidate, destination)
                        shutil.copy2(rendered, destination.with_suffix(".verification.pdf"))
                        break
                    rendered.unlink()
                if not verified:
                    if actual_pages is not None:
                        raise ValueError(
                            f"تعذر الحفاظ على عدد الصفحات في Word: المتوقع {len(pages)}، والناتج {actual_pages}. اختصر النص أو راجع توزيع الصفحة."
                        )
                    shutil.copy2(candidate, destination)
    return _export_report(
        library,
        book,
        pages,
        selection,
        start_page,
        end_page,
        destination,
        "word",
        verified=verified,
        actual_pages=actual_pages,
        font_scale=used_scale,
    )


def _markdown_from_html(value: str) -> str:
    blocks = rich_html_blocks(value)
    rendered: list[str] = []
    for block in blocks:
        if block.get("separator"):
            rendered.append("---")
            continue
        parts: list[str] = []
        list_kind = None
        for item in block.get("runs", []):
            text = str(item.get("text", ""))
            text = re.sub(r"([\\`*_\[\]<>])", r"\\\1", text)
            if item.get("bold"):
                text = f"**{text}**"
            if item.get("italic"):
                text = f"*{text}*"
            if item.get("underline"):
                text = f"<u>{text}</u>"
            if item.get("superscript"):
                text = f"<sup>{text}</sup>"
            if item.get("subscript"):
                text = f"<sub>{text}</sub>"
            parts.append(text)
            list_kind = list_kind or item.get("list")
        text = "".join(parts).strip()
        if not text:
            continue
        if block.get("heading_level"):
            text = f"{'#' * int(block['heading_level'])} {text}"
        elif block.get("quote"):
            text = "\n".join(f"> {line}" for line in text.splitlines())
        elif list_kind:
            text = f"{'1.' if list_kind == 'ordered' else '-'} {text}"
        rendered.append(text)
    return "\n\n".join(rendered)


def _page_markdown(page: dict[str, Any]) -> str:
    if page.get("content_format") == "markdown" and page.get("content_markdown"):
        return str(page["content_markdown"]).strip()
    return _markdown_from_html(page.get("content_html", ""))


def export_markdown(
    library: Library,
    book_id: str,
    start_page: int,
    end_page: int,
    destination: Path,
    *,
    include_unreviewed: bool = False,
) -> dict[str, Any]:
    book, pages, selection = _export_selection(
        library, book_id, start_page, end_page, include_unreviewed
    )
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    documents = [f"# {book['name']}"]
    for page in pages:
        documents.append(_page_markdown(page))
    destination.write_text("\n\n\n".join(documents) + "\n", encoding="utf-8")
    return _export_report(
        library,
        book,
        pages,
        selection,
        start_page,
        end_page,
        destination,
        "markdown",
        verified=True,
        actual_pages=len(pages),
    )


def _html_document(book: dict[str, Any], pages: list[dict[str, Any]]) -> str:
    title = html.escape(str(book["name"]))
    articles: list[str] = []
    for index, page in enumerate(pages):
        source_number = int(page["number"])
        printed_page = html.escape(str(page.get("printed_page") or source_number))
        content = sanitize_rich_html(page.get("content_html", ""))
        hidden = "" if index == 0 else " hidden"
        width = max(float(page.get("width") or 420), 1)
        height = max(float(page.get("height") or 594), 1)
        direction = page.get("text_direction", "auto")
        if direction not in {"rtl", "ltr"}:
            direction = "auto"
        articles.append(
            f'<article class="book-page" data-page="{source_number}" '
            f'data-printed-page="{printed_page}" style="--page-ratio:{width}/{height}"{hidden}>'
            f'<div class="page-content" dir="{direction}">{content}</div>'
            f'<footer class="printed-page">{printed_page}</footer>'
            "</article>"
        )
    pages_html = "\n".join(articles)
    return f'''<!doctype html>
<html lang="ar" dir="rtl">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title>
<style>
:root {{ color-scheme: light; font-family: "Noto Naskh Arabic", "Geeza Pro", Arial, sans-serif; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; min-height: 100vh; background: #eceaf2; color: #1d1b24; }}
button, input {{ font: inherit; }}
.book-toolbar {{ position: sticky; top: 0; z-index: 5; display: grid; grid-template-columns: auto minmax(150px, 1fr) auto; gap: 12px; align-items: center; padding: 12px clamp(14px, 3vw, 32px); background: rgba(255,255,255,.96); border-bottom: 1px solid #d8d4e0; box-shadow: 0 4px 18px rgba(40,30,60,.08); }}
.navigation {{ display: flex; gap: 8px; align-items: center; }}
.navigation button, .page-jump button {{ min-height: 42px; padding: 7px 15px; border: 1px solid #bdb6ca; border-radius: 9px; background: #fff; color: #3d3154; cursor: pointer; }}
.navigation button:hover, .page-jump button:hover {{ background: #f1edf7; }}
.navigation button:disabled {{ cursor: default; opacity: .45; }}
.page-jump {{ display: flex; gap: 7px; align-items: center; white-space: nowrap; }}
.page-jump input {{ width: 86px; min-height: 42px; border: 1px solid #bdb6ca; border-radius: 9px; padding: 6px 9px; text-align: center; }}
.search {{ position: relative; width: min(460px, 100%); justify-self: center; }}
.search input {{ width: 100%; min-height: 42px; border: 1px solid #bdb6ca; border-radius: 9px; padding: 7px 13px; background: #fff; }}
#search-results {{ position: absolute; inset-inline: 0; top: calc(100% + 7px); max-height: min(420px, 62vh); overflow: auto; margin: 0; padding: 6px; list-style: none; border: 1px solid #d8d4e0; border-radius: 10px; background: #fff; box-shadow: 0 12px 32px rgba(40,30,60,.16); }}
#search-results:empty {{ display: none; }}
#search-results button {{ width: 100%; border: 0; border-radius: 7px; padding: 9px 10px; background: transparent; color: inherit; text-align: right; cursor: pointer; }}
#search-results button:hover {{ background: #f1edf7; }}
#search-results strong {{ color: #684797; }}
.book-stage {{ display: grid; place-items: start center; padding: clamp(18px, 4vw, 42px); }}
.book-page {{ position: relative; width: min(820px, 94vw); aspect-ratio: var(--page-ratio); min-height: 72vh; padding: clamp(34px, 7vw, 72px); overflow: auto; background: #fff; border: 1px solid #d2ced7; box-shadow: 0 15px 40px rgba(35,29,45,.13); }}
.book-page[hidden] {{ display: none; }}
.page-content {{ text-align: start; font-size: clamp(18px, 2.25vw, 23px); line-height: 1.75; }}
.page-content p, .page-content h1, .page-content h2, .page-content h3, .page-content h4, .page-content h5, .page-content h6, .page-content blockquote {{ margin: 0 0 .32em; }}
.page-content h1, .page-content h2, .page-content h3, .page-content h4, .page-content h5, .page-content h6 {{ text-align: center; line-height: 1.45; }}
.page-content blockquote, .page-content .ql-size-small {{ margin-inline: 0; font-size: .82em; line-height: 1.65; }}
.page-content blockquote {{ padding-inline: 1.2em; }}
.page-content hr {{ border: 0; border-top: 1px solid #777; margin: 1.1em 0 .8em; }}
.page-content .ql-align-center {{ text-align: center; }}
.page-content .ql-align-right {{ text-align: right; }}
.page-content .ql-align-left {{ text-align: left; }}
.page-content .ql-align-justify {{ text-align: justify; }}
.page-content .ql-size-large {{ font-size: 1.35em; }}
.page-content .ql-size-huge {{ font-size: 1.75em; }}
.page-content .ql-font-arial {{ font-family: Arial, sans-serif; }}
.page-content .ql-font-plex {{ font-family: "IBM Plex Sans Arabic", Arial, sans-serif; }}
.printed-page {{ position: absolute; bottom: 18px; inset-inline: 0; text-align: center; font-size: 14px; color: #615c68; }}
.status {{ color: #5d5667; font-size: 14px; }}
@media (max-width: 760px) {{ .book-toolbar {{ grid-template-columns: 1fr; }} .navigation, .page-jump {{ justify-content: center; }} .search {{ grid-row: 1; }} .book-page {{ padding: 34px 24px 56px; }} }}
@media print {{ body {{ background: #fff; }} .book-toolbar {{ display: none; }} .book-stage {{ display: block; padding: 0; }} .book-page, .book-page[hidden] {{ display: block; width: 100%; min-height: 0; box-shadow: none; border: 0; break-after: page; }} }}
</style>
</head>
<body>
<header class="book-toolbar">
  <nav class="navigation" aria-label="التنقل بين الصفحات">
    <button id="previous-page" type="button">السابق</button>
    <button id="next-page" type="button">التالي</button>
    <span id="page-status" class="status" aria-live="polite"></span>
  </nav>
  <div class="search">
    <input id="book-search" type="search" placeholder="ابحث داخل الكتاب" autocomplete="off" aria-label="البحث داخل الكتاب">
    <ol id="search-results" aria-label="نتائج البحث"></ol>
  </div>
  <form id="page-jump" class="page-jump">
    <label for="page-number">صفحة</label>
    <input id="page-number" type="number" inputmode="numeric" aria-label="رقم الصفحة">
    <button type="submit">انتقال</button>
  </form>
</header>
<main class="book-stage" aria-label="{title}">
{pages_html}
</main>
<script>
(() => {{
  "use strict";
  const pages = [...document.querySelectorAll(".book-page")];
  const previous = document.getElementById("previous-page");
  const next = document.getElementById("next-page");
  const status = document.getElementById("page-status");
  const pageNumber = document.getElementById("page-number");
  const pageJump = document.getElementById("page-jump");
  const search = document.getElementById("book-search");
  const results = document.getElementById("search-results");
  let current = 0;
  const normalize = value => String(value || "").normalize("NFKD").replace(/[\u064B-\u065F\u0670\u06D6-\u06ED\u0640]/g, "").toLocaleLowerCase("ar");
  const show = index => {{
    current = Math.max(0, Math.min(index, pages.length - 1));
    pages.forEach((page, item) => page.hidden = item !== current);
    const sourcePage = pages[current].dataset.page;
    pageNumber.value = sourcePage;
    status.textContent = `الصفحة ${{current + 1}} من ${{pages.length}}`;
    previous.disabled = current === 0;
    next.disabled = current === pages.length - 1;
    history.replaceState(null, "", `#page-${{sourcePage}}`);
    window.scrollTo({{ top: 0, behavior: "smooth" }});
  }};
  const goToSourcePage = value => {{
    const index = pages.findIndex(page => Number(page.dataset.page) === Number(value));
    if (index >= 0) show(index);
    else pageNumber.setCustomValidity("هذه الصفحة غير موجودة في الملف المصدّر");
  }};
  previous.addEventListener("click", () => show(current - 1));
  next.addEventListener("click", () => show(current + 1));
  pageJump.addEventListener("submit", event => {{
    event.preventDefault();
    pageNumber.setCustomValidity("");
    goToSourcePage(pageNumber.value);
    pageNumber.reportValidity();
  }});
  pageNumber.addEventListener("input", () => pageNumber.setCustomValidity(""));
  document.addEventListener("keydown", event => {{
    if (event.target.matches("input")) return;
    if (event.key === "ArrowRight") show(current + 1);
    if (event.key === "ArrowLeft") show(current - 1);
  }});
  search.addEventListener("input", () => {{
    results.replaceChildren();
    const query = normalize(search.value.trim());
    if (!query) return;
    pages.forEach((page, index) => {{
      const plain = page.querySelector(".page-content").textContent.replace(/\\s+/g, " ").trim();
      const normalized = normalize(plain);
      const match = normalized.indexOf(query);
      if (match < 0) return;
      const start = Math.max(0, match - 45);
      const snippet = plain.slice(start, Math.min(plain.length, match + search.value.length + 70));
      const item = document.createElement("li");
      const button = document.createElement("button");
      button.type = "button";
      const label = document.createElement("strong");
      label.textContent = `صفحة ${{page.dataset.page}}: `;
      button.append(label, document.createTextNode(snippet));
      button.addEventListener("click", () => {{ show(index); results.replaceChildren(); }});
      item.append(button);
      results.append(item);
    }});
    if (!results.children.length) {{
      const item = document.createElement("li");
      item.textContent = "لا توجد نتائج";
      item.style.padding = "9px 10px";
      results.append(item);
    }}
  }});
  const initialPage = Number(location.hash.replace("#page-", ""));
  const initialIndex = pages.findIndex(page => Number(page.dataset.page) === initialPage);
  show(initialIndex >= 0 ? initialIndex : 0);
}})();
</script>
</body>
</html>
'''


def export_html(
    library: Library,
    book_id: str,
    start_page: int,
    end_page: int,
    destination: Path,
    *,
    include_unreviewed: bool = False,
) -> dict[str, Any]:
    book, pages, selection = _export_selection(
        library, book_id, start_page, end_page, include_unreviewed
    )
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(_html_document(book, pages), encoding="utf-8")
    return _export_report(
        library,
        book,
        pages,
        selection,
        start_page,
        end_page,
        destination,
        "html",
        verified=True,
        actual_pages=len(pages),
    )
